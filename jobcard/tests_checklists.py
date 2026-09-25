"""Checklists module and the work order Checklist section."""
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from Inventory.models import Department, Equipment, EquipmentDescription
from jobcard.checklists import ChecklistIncomplete, read_answers
from jobcard.models import ChecklistItem, ChecklistTemplate, WorkOrderChecklistEntry, jobcard
from users.models import UserProfile
from workshop.models import Workshop

from core.tests.excel_helpers import ExcelImportTestMixin

from .tests import SIGNATURE

User = get_user_model()


class ChecklistTests(TestCase):
    def setUp(self):
        self.workshop = Workshop.objects.create(name="Biomed", category="maintenance")
        self.other_workshop = Workshop.objects.create(name="Annex", category="maintenance")
        self.department = Department.objects.create(name="ICU", workshop=self.workshop)
        self.description = EquipmentDescription.objects.create(name="Defibrillator")
        self.equipment = Equipment.objects.create(description=self.description, model="D1", serial_number="DEF-1",
                                                  department=self.department, status="Working")
        self.tech = self._user("cl_tech", "Tech", workshop=self.workshop, level="Engineer")
        self.outsider = self._user("cl_out", "Tech", workshop=self.other_workshop, level="Engineer")

        self.template = ChecklistTemplate.objects.create(
            equipment_description=self.description, task_type="PPM", title="Defib PPM",
            instructions="Disconnect from patient.")
        self.visual = ChecklistItem.objects.create(template=self.template, order=1, task="Inspect paddles")
        self.energy = ChecklistItem.objects.create(template=self.template, order=2, task="Discharge test",
                                                   response_type="value", expected_result="200 J ± 15%")
        self.optional = ChecklistItem.objects.create(template=self.template, order=3, task="Clean case",
                                                     is_required=False)
        # Applies to every task on this description
        self.any_template = ChecklistTemplate.objects.create(
            equipment_description=self.description, task_type="Any", title="Electrical safety")
        self.leakage = ChecklistItem.objects.create(template=self.any_template, order=1, task="Leakage current",
                                                    response_type="pass_fail")

    def _user(self, username, role, **profile):
        user = User.objects.create_user(username=username, password="pw12345!")
        UserProfile.objects.update_or_create(user=user, defaults={
            "role": role, "must_change_password": False, "has_uploaded_signature": True, **profile})
        return user

    def _answers(self, **overrides):
        data = {
            "checklist_template": str(self.template.id),
            f"checklist_result_{self.visual.id}": "done",
            f"checklist_result_{self.energy.id}": "pass",
            f"checklist_value_{self.energy.id}": "198 J",
        }
        data.update(overrides)
        return data

    def _post(self, data):
        """A QueryDict, as the view receives it (read_answers uses getlist)."""
        from django.http import QueryDict
        q = QueryDict(mutable=True)
        for key, value in data.items():
            if isinstance(value, list):
                q.setlist(key, value)
            else:
                q[key] = value
        return q

    def _raise(self, action="PPM", **extra):
        self.client.force_login(self.tech)
        return self.client.post(reverse("jobcard:create_job_card"), {
            "priority_level": "Medium", "department": self.department.id, "equipment": self.equipment.id,
            "job_description": "Quarterly PPM", "action_taken": action, "time_started": "09:00",
            "signature_data": SIGNATURE, "spare_parts_data": "[]", "labor_cost": "0", "additional_costs": "0",
            **extra,
        })

    # ── rules ────────────────────────────────────────────────────────────────

    def test_checklists_suggested_by_description_and_task(self):
        titles = [t.title for t in ChecklistTemplate.for_work(self.equipment, "PPM")]
        self.assertEqual(set(titles), {"Defib PPM", "Electrical safety"})
        titles = {t.title for t in ChecklistTemplate.for_work(self.equipment, "Repair")}
        self.assertEqual(titles, {"Electrical safety"})

    def test_checklist_for_all_equipment_is_suggested_after_specific_ones(self):
        ChecklistTemplate.objects.create(equipment_description=None, task_type="Any", title="AAA general")
        titles = [t.title for t in ChecklistTemplate.for_work(self.equipment, "PPM")]
        self.assertEqual(titles[-1], "AAA general")

    def test_complete_answers_pass(self):
        template, answers = read_answers(self._post(self._answers()))
        self.assertEqual(template, self.template)
        self.assertEqual(len(answers), 2)  # the unanswered optional step is not recorded

    def test_no_checklist_selected_is_allowed(self):
        template, answers = read_answers(self._post({}))
        self.assertIsNone(template)
        self.assertEqual(answers, [])

    def test_required_step_missing(self):
        post = self._answers()
        del post[f"checklist_result_{self.visual.id}"]
        with self.assertRaisesMessage(ChecklistIncomplete, "Inspect paddles"):
            read_answers(self._post(post))

    def test_value_step_needs_reading(self):
        with self.assertRaisesMessage(ChecklistIncomplete, "enter the reading"):
            read_answers(self._post(self._answers(**{f"checklist_value_{self.energy.id}": ""})))

    def test_failure_needs_a_note(self):
        post = self._answers(**{f"checklist_result_{self.energy.id}": "fail"})
        with self.assertRaisesMessage(ChecklistIncomplete, "say why"):
            read_answers(self._post(post))

    def test_answer_must_fit_the_step(self):
        with self.assertRaisesMessage(ChecklistIncomplete, "invalid answer"):
            read_answers(self._post(self._answers(**{f"checklist_result_{self.visual.id}": "pass"})))

    def test_disabled_checklist_cannot_be_selected(self):
        self.template.active_status = False
        self.template.save()
        with self.assertRaisesMessage(ChecklistIncomplete, "no longer available"):
            read_answers(self._post(self._answers()))

    def test_custom_steps(self):
        post = self._post({"custom_keys": ["a", "b", "c"],
                           "custom_task_a": "Check battery", "custom_result_a": "done",
                           "custom_task_b": "Replace mains fuse", "custom_result_b": "not_done",
                           "custom_note_b": "no spare in store",
                           "custom_task_c": "", "custom_result_c": ""})  # blank row ignored
        template, answers = read_answers(post)
        self.assertIsNone(template)
        self.assertEqual([a["task"] for a in answers], ["Check battery", "Replace mains fuse"])
        self.assertTrue(all(a["item"] is None for a in answers))

    def test_custom_step_needs_a_result(self):
        with self.assertRaisesMessage(ChecklistIncomplete, "mark it done"):
            read_answers(self._post({"custom_keys": ["a"], "custom_task_a": "Check battery"}))

    def test_answer_time_from_the_browser_is_kept_when_plausible(self):
        from django.utils import timezone
        from datetime import timedelta
        earlier = timezone.now() - timedelta(minutes=40)
        far = timezone.now() - timedelta(days=30)
        _, answers = read_answers(self._post(self._answers(**{
            f"checklist_at_{self.visual.id}": earlier.isoformat(),
            f"checklist_at_{self.energy.id}": far.isoformat(),
        })))
        by_task = {a["task"]: a["completed_at"] for a in answers}
        self.assertLess(abs((by_task["Inspect paddles"] - earlier).total_seconds()), 1)
        self.assertLess((timezone.now() - by_task["Discharge test"]).total_seconds(), 60)  # implausible -> now

    # ── through the work order form ──────────────────────────────────────────

    def test_work_order_stores_answers_with_who_and_when(self):
        self._raise(**self._answers(**{
            f"checklist_result_{self.energy.id}": "fail",
            f"checklist_note_{self.energy.id}": "150 J, below limit",
            "custom_keys": "k1", "custom_task_k1": "Check battery", "custom_result_k1": "done",
        }))
        card = jobcard.objects.get(equipment=self.equipment)
        entries = {e.task: e for e in card.checklist_entries.all()}
        self.assertEqual(set(entries), {"Inspect paddles", "Discharge test", "Check battery"})
        self.assertEqual(entries["Discharge test"].value, "198 J")
        self.assertTrue(entries["Discharge test"].is_problem)
        self.assertEqual(entries["Inspect paddles"].template, self.template)
        self.assertEqual(entries["Inspect paddles"].completed_by, self.tech)
        self.assertIsNotNone(entries["Inspect paddles"].completed_at)
        self.assertIsNone(entries["Check battery"].template)
        self.assertTrue(entries["Check battery"].is_custom)

        # Editing the template later does not rewrite history.
        self.visual.task = "Inspect paddles and cables"
        self.visual.save()
        self.assertTrue(card.checklist_entries.filter(task="Inspect paddles").exists())

    def test_work_order_refused_until_required_steps_answered(self):
        response = self._raise(checklist_template=str(self.template.id))
        self.assertEqual(response.status_code, 200)  # form re-rendered
        self.assertFalse(jobcard.objects.filter(equipment=self.equipment).exists())
        self.assertContains(response, "Checklist incomplete")

    def test_work_order_without_checklist(self):
        self._raise()
        self.assertTrue(jobcard.objects.filter(equipment=self.equipment).exists())

    def test_pdf_includes_checklist(self):
        from jobcard.modern_jobcard_pdf import generate_jobcard_pdf
        self._raise(**self._answers())
        card = jobcard.objects.get(equipment=self.equipment)
        self.assertTrue(generate_jobcard_pdf(card).getvalue().startswith(b"%PDF"))

    def test_work_order_detail_shows_checklist(self):
        self._raise(**self._answers(**{f"checklist_note_{self.visual.id}": "new paddles fitted"}))
        card = jobcard.objects.get(equipment=self.equipment)
        response = self.client.get(reverse("jobcard:work_order_detail", args=[card.id]))
        self.assertContains(response, card.work_order_number)
        self.assertContains(response, "Defib PPM")
        self.assertContains(response, "new paddles fitted")
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("jobcard:work_order_detail", args=[card.id])).status_code, 404)

    def test_lookup_endpoint(self):
        self.client.force_login(self.tech)
        url = reverse("jobcard:checklist_for_work")
        other = ChecklistTemplate.objects.create(
            equipment_description=EquipmentDescription.objects.create(name="Ventilator"), title="Vent PPM")
        data = self.client.get(url, {"equipment_id": self.equipment.id, "action": "PPM"}).json()
        self.assertEqual({t["title"] for t in data["suggested"]}, {"Defib PPM", "Electrical safety"})
        self.assertIn(str(other.id), [t["id"] for t in data["others"]])
        data = self.client.get(url, {"template_id": self.template.id}).json()
        self.assertEqual([i["task"] for i in data["template"]["items"]],
                         ["Inspect paddles", "Discharge test", "Clean case"])

    def test_lookup_is_scoped_to_the_users_workshop(self):
        self.client.force_login(self.outsider)
        response = self.client.get(reverse("jobcard:checklist_for_work"),
                                   {"equipment_id": self.equipment.id, "action": "PPM"})
        self.assertEqual(response.status_code, 404)

    # ── settings module ──────────────────────────────────────────────────────

    def test_create_checklist_with_steps(self):
        self.client.force_login(self.tech)
        self.client.post(reverse("jobcard:checklist_create"), {
            "title": "Defib repair checks", "equipment_description": self.description.id, "task_type": "Repair",
            "instructions": "Replace fuses only with rated parts.",
            "item_id": ["", ""], "item_task": ["Check fuse", "Self test"],
            "item_guidance": ["Open rear panel", ""], "item_expected": ["", "PASS on display"],
            "item_response_type": ["check", "pass_fail"], "item_required": ["1", "0"],
        })
        template = ChecklistTemplate.objects.get(title="Defib repair checks")
        items = list(template.items.order_by("order"))
        self.assertEqual([i.task for i in items], ["Check fuse", "Self test"])
        self.assertFalse(items[1].is_required)

    def test_removed_step_is_deactivated_not_deleted(self):
        self.client.force_login(self.tech)
        self.client.post(reverse("jobcard:checklist_edit", args=[self.template.id]), {
            "title": "Defib PPM", "equipment_description": self.description.id, "task_type": "PPM",
            "item_id": [str(self.energy.id)], "item_task": ["Discharge test"], "item_guidance": [""],
            "item_expected": ["200 J"], "item_response_type": ["value"], "item_required": ["1"],
        })
        self.visual.refresh_from_db()
        self.energy.refresh_from_db()
        self.assertFalse(self.visual.active_status)
        self.assertEqual(self.energy.order, 1)
        self.assertEqual(self.energy.expected_result, "200 J")

    def test_settings_pages_render(self):
        self.client.force_login(self.tech)
        self.assertEqual(self.client.get(reverse("jobcard:checklist_settings")).status_code, 200)
        self.assertEqual(self.client.get(reverse("jobcard:checklist_create")).status_code, 200)
        self.assertEqual(self.client.get(reverse("jobcard:checklist_edit", args=[self.template.id])).status_code, 200)

    def test_checklist_for_all_equipment(self):
        self.client.force_login(self.tech)
        self.client.post(reverse("jobcard:checklist_create"), {
            "title": "General safety", "equipment_description": "", "task_type": "Any",
            "item_id": [""], "item_task": ["Visual inspection"], "item_guidance": [""], "item_expected": [""],
            "item_response_type": ["check"], "item_required": ["1"],
        })
        template = ChecklistTemplate.objects.get(title="General safety")
        self.assertIsNone(template.equipment_description)
        self.assertEqual(template.applies_to, "All equipment")

    def test_list_shows_counts_and_links_to_detail(self):
        self._raise(**self._answers())
        response = self.client.get(reverse("jobcard:checklist_settings"))
        self.assertContains(response, reverse("jobcard:checklist_detail", args=[self.template.id]))
        row = next(t for t in response.context["templates"] if t.id == self.template.id)
        self.assertEqual((row.item_count, row.use_count), (3, 1))

    def test_detail_page(self):
        self._raise(**self._answers())
        card = jobcard.objects.get(equipment=self.equipment)
        response = self.client.get(reverse("jobcard:checklist_detail", args=[self.template.id]))
        self.assertContains(response, "Disconnect from patient.")
        self.assertContains(response, "200 J ± 15%")
        self.assertContains(response, card.work_order_number)
        self.assertEqual(response.context["template"].use_count, 1)

    def test_checklists_are_for_technologists_only(self):
        nurse = self._user("cl_nic", "NIC", department=self.department)
        self.client.force_login(nurse)
        for name, args in (("checklist_settings", []), ("checklist_detail", [self.template.id]),
                           ("checklist_edit", [self.template.id])):
            self.assertEqual(self.client.get(reverse(f"jobcard:{name}", args=args)).status_code, 302, name)

    # ── starters and Machine Reports history ─────────────────────────────────

    def test_copy_a_starter_checklist(self):
        from jobcard.starter_checklists import STARTERS
        monitor = EquipmentDescription.objects.create(name="Patient Monitor")
        self.client.force_login(self.tech)
        response = self.client.post(reverse("jobcard:checklist_from_starter"),
                                    {"starter": "patient_monitor", "equipment_description": monitor.id})
        template = ChecklistTemplate.objects.get(equipment_description=monitor)
        self.assertRedirects(response, reverse("jobcard:checklist_edit", args=[template.id]))
        tasks = list(template.items.order_by("order").values_list("task", flat=True))
        self.assertEqual(len(tasks), len(STARTERS["patient_monitor"]["items"]))
        self.assertTrue(any("NIBP cuff" in t for t in tasks))

    def test_machine_reports_history_shows_checklists(self):
        self._raise(**self._answers(**{f"checklist_note_{self.visual.id}": "new paddles"}))
        card = jobcard.objects.get(equipment=self.equipment)
        jobcard.objects.filter(pk=card.pk).update(status="Approved")
        self.client.force_login(self.tech)
        data = self.client.get(reverse("equipment_repair_details", args=[self.equipment.id])).json()
        self.assertEqual(data["work_history"][0]["action"], "PPM")
        tasks = {t["task"]: t for t in data["task_last_done"]}
        self.assertEqual(tasks["Discharge test"]["value"], "198 J")
        self.assertEqual(tasks["Inspect paddles"]["work_order"], str(card.id)[:8].upper())

    def test_history_export_has_checklist_sheets(self):
        import io

        from openpyxl import load_workbook
        self._raise(**self._answers())
        jobcard.objects.filter(equipment=self.equipment).update(status="Approved")
        self.client.force_login(self.tech)
        response = self.client.get(reverse("export_equipment_history", args=[self.equipment.id]))
        wb = load_workbook(io.BytesIO(response.content))
        self.assertIn("Checklist Tasks Last Done", wb.sheetnames)
        self.assertEqual(wb["Checklist History"].max_row, 3)  # header + 2 answered steps


class ChecklistImportTests(ExcelImportTestMixin, TestCase):
    HEADERS = ["Checklist Name", "Applies To", "Task Type", "Description", "Step No.", "Step", "Instructions",
               "Expected Result", "Response Type", "Required"]

    def setUp(self):
        super().setUp()
        self.workshop = Workshop.objects.create(name="Biomed", category="maintenance")
        self.description = EquipmentDescription.objects.create(name="Patient Monitor")
        self.tech = self.make_user("imp_tech", "Tech", workshop=self.workshop, level="Engineer")
        self.client.force_login(self.tech)

    def test_template_download(self):
        from core.tests.excel_helpers import read_response_workbook
        wb = read_response_workbook(self.client.get(reverse("jobcard:checklist_import_template")))
        self.assertEqual([c.value for c in wb["Checklists"][1]], self.HEADERS)

    def test_steps_group_into_checklists(self):
        rows = [self.HEADERS,
                ["Monitor PM", "patient monitor", "PPM", "Quarterly", 2, "Verify alarms", "", "Alarm sounds",
                 "Pass / Fail", "Yes"],
                ["Monitor PM", "Patient Monitor", "PPM", "", 1, "Inspect power cable", "", "", "", ""],
                ["General check", "", "", "Any device", "", "Clean equipment", "", "", "", "No"],
                ["Vent PM", "Ventilator", "PPM", "", "", "Check O2 cell", "", "", "Reading", ""],
                ["Broken", "", "Weekly", "", "", "Step", "", "", "", ""]]
        url = reverse("jobcard:checklist_import_upload")

        preview = self.post_workbook(url, {"Checklists": rows}).json()
        self.assertEqual(preview["counts"], {"create": 3, "skip": 0, "error": 2})  # Ventilator missing, bad task
        self.assertFalse(ChecklistTemplate.objects.exists())

        data = self.post_workbook(url, {"Checklists": rows}, commit=True, create_missing="true").json()
        self.assertEqual(data["counts"], {"create": 4, "skip": 0, "error": 1})
        self.assertEqual(data["created"], {"Equipment descriptions": ["Ventilator"]})
        monitor = ChecklistTemplate.objects.get(title="Monitor PM")
        self.assertEqual((monitor.equipment_description, monitor.task_type, monitor.instructions),
                         (self.description, "PPM", "Quarterly"))
        self.assertEqual([(i.task, i.response_type) for i in monitor.active_items()],
                         [("Inspect power cable", "check"), ("Verify alarms", "pass_fail")])
        general = ChecklistTemplate.objects.get(title="General check")
        self.assertIsNone(general.equipment_description)
        self.assertFalse(general.items.get().is_required)
        self.assertEqual(ChecklistTemplate.objects.get(title="Vent PM").items.get().response_type, "value")

        again = self.post_workbook(url, {"Checklists": rows[:3]}).json()
        self.assertEqual(again["counts"], {"create": 0, "skip": 2, "error": 0})
