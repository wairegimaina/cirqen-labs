"""Bulk equipment import from Excel.

Two endpoints:

* :func:`download_equipment_import_template` – hands the user a pre-formatted
  ``.xlsx`` with the expected columns, dropdowns for Department/Status and a
  reference sheet listing the values that already exist in their workshop.
* :func:`upload_equipment_excel` – parses an uploaded workbook, validates every
  row and either previews the outcome (``commit`` absent/false) or writes it.

The preview and the commit run the *same* code path inside a transaction; the
preview simply rolls back at the end, so what the user is shown is exactly what
they get. Rows are saved through ``Equipment.save()`` (not ``bulk_create``) so
model validation, category inheritance and ``needs_sync`` behave as they do in
:func:`Inventory.views.equipment.add_inventory`.

Uploads require HQ. They are refused while HQ is unreachable, and a commit is
pushed to HQ's sync API *inside* the local transaction: unless HQ accepts every
row, the local transaction is rolled back and nothing is saved.
"""
import logging
import re

from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core import hq_link
from Inventory.models import Department, Equipment, EquipmentDescription, Manufacturer

logger = logging.getLogger(__name__)

# Column order matches export_equipment_to_excel so an exported file can be
# edited and uploaded straight back.
COLUMNS = ["description", "manufacturer", "model", "serial_number", "department", "status",
           # Optional, after the required six so older files still line up.
           "asset_tag"]

COLUMN_LABELS = {
    "description": "Description",
    "manufacturer": "Manufacturer",
    "model": "Model",
    "serial_number": "Serial Number",
    "department": "Department",
    "status": "Status",
    "asset_tag": "Hospital Asset No.",
}

# Accepted header spellings, keyed by normalized text (lowercase, alphanumeric).
HEADER_ALIASES = {
    "description": "description",
    "equipmentdescription": "description",
    "equipment": "description",
    "equipmentname": "description",
    "item": "description",
    "manufacturer": "manufacturer",
    "make": "manufacturer",
    "brand": "manufacturer",
    "model": "model",
    "modelno": "model",
    "modelnumber": "model",
    "serialnumber": "serial_number",
    "serialno": "serial_number",
    "serial": "serial_number",
    "sn": "serial_number",
    "department": "department",
    "dept": "department",
    "unit": "department",
    "location": "department",
    "status": "status",
    "condition": "status",
    "hospitalassetno": "asset_tag",
    "assetno": "asset_tag",
    "assetnumber": "asset_tag",
    "assettag": "asset_tag",
    "tagno": "asset_tag",
    "barcode": "asset_tag",
}

STATUS_CHOICES = ["Working", "Not working", "Under repair"]
STATUS_ALIASES = {
    "working": "Working",
    "functional": "Working",
    "operational": "Working",
    "ok": "Working",
    "good": "Working",
    "notworking": "Not working",
    "nonfunctional": "Not working",
    "faulty": "Not working",
    "broken": "Not working",
    "down": "Not working",
    "underrepair": "Under repair",
    "repair": "Under repair",
    "inrepair": "Under repair",
    "servicing": "Under repair",
}

# Field lengths mirrored from the models so users get a clean row-level error
# instead of a database exception.
MAX_LENGTHS = {
    "description": 200,
    "manufacturer": 150,
    "model": 100,
    "serial_number": 100,
    "asset_tag": 100,
}

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 5000                   # refuse absurd files outright
MAX_HEADER_SCAN = 15              # rows to search for the header row
MAX_REPORT_ROWS = 500             # per-row detail returned to the browser


HQ_OFFLINE_MESSAGE = (
    "Bulk upload needs a connection to the HQ server, and HQ is not reachable "
    "right now. Try again once this machine is back online."
)


class _DryRunRollback(Exception):
    """Raised to unwind the transaction after a preview run."""


class _HQRejected(Exception):
    """Raised to unwind a commit that HQ did not fully accept."""


def _normalize_header(value):
    """``"Serial No."`` -> ``"serialno"``."""
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def _normalize_status(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _cell_text(value):
    """Coerce a cell to trimmed text, without turning numbers into floats."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _import_scope(request):
    """Resolve who may import and into which workshop.

    Returns ``(profile, workshop, error_message)``. Only Technologists import,
    matching the restriction already enforced on equipment-description creation.
    """
    profile = getattr(request.user, "userprofile", None)
    if profile is None:
        return None, None, "User profile not found."
    if profile.role != "Tech":
        return None, None, "Only Technologists can bulk-upload equipment."
    if not profile.workshop:
        return None, None, "Your profile is not associated with a workshop."
    return profile, profile.workshop, None


@login_required
@require_http_methods(["GET"])
def download_equipment_import_template(request):
    """Build the blank import workbook, pre-loaded with the user's own values."""
    profile, workshop, error = _import_scope(request)
    if error:
        return JsonResponse({"success": False, "error": error}, status=403)

    departments = list(
        Department.objects.filter(workshop=workshop, active_status=True)
        .order_by("name")
        .values_list("name", flat=True)
    )
    descriptions = list(
        EquipmentDescription.objects.filter(active_status=True)
        .order_by("name")
        .values_list("name", flat=True)
    )
    manufacturers = list(
        Manufacturer.objects.filter(active_status=True)
        .order_by("name")
        .values_list("name", flat=True)
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Equipment"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, key in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=COLUMN_LABELS[key])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 26

    ws.freeze_panes = "A2"

    # Reference sheet: feeds the dropdowns and shows what already exists.
    ref = wb.create_sheet("Reference")
    ref.column_dimensions["A"].width = 34
    ref.column_dimensions["B"].width = 34
    ref.column_dimensions["C"].width = 34
    ref.column_dimensions["D"].width = 20

    for col, (title, values) in enumerate(
        [
            ("Departments", departments),
            ("Equipment Descriptions", descriptions),
            ("Manufacturers", manufacturers),
            ("Status", STATUS_CHOICES),
        ],
        start=1,
    ):
        ref.cell(row=1, column=col, value=title).font = Font(bold=True)
        for offset, value in enumerate(values, start=2):
            ref.cell(row=offset, column=col, value=value)

    # Dropdowns for the two closed-vocabulary columns.
    if departments:
        dept_range = f"'Reference'!$A$2:$A${len(departments) + 1}"
        dv_dept = DataValidation(type="list", formula1=dept_range, allow_blank=True)
        ws.add_data_validation(dv_dept)
        dv_dept.add(f"E2:E{MAX_ROWS + 1}")

    status_range = f"'Reference'!$D$2:$D${len(STATUS_CHOICES) + 1}"
    dv_status = DataValidation(type="list", formula1=status_range, allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"F2:F{MAX_ROWS + 1}")

    # Notes sheet rather than instruction rows in the data sheet, so the parser
    # never has to guess where the data starts.
    notes = wb.create_sheet("Instructions")
    notes.column_dimensions["A"].width = 110
    guidance = [
        "How to use this template",
        "",
        "1. Enter one piece of equipment per row on the 'Equipment' sheet, starting at row 2.",
        "2. Do not rename, reorder or delete the header row.",
        "3. Description, Model, Serial Number, Department and Status are required. Manufacturer is optional.",
        "4. Serial Number must be unique. It is stored in UPPERCASE and duplicates are reported, not imported.",
        "5. Department must already exist in your workshop - see the 'Reference' sheet for valid names.",
        "6. Status must be one of: Working, Not working, Under repair.",
        "7. Descriptions and Manufacturers that do not exist yet can be created automatically during upload",
        "   (tick 'Create missing descriptions and manufacturers' in the upload dialog).",
        f"8. Maximum {MAX_ROWS} rows per upload.",
        "",
        "You will always see a preview of what will be imported before anything is saved.",
    ]
    for row, line in enumerate(guidance, start=1):
        cell = notes.cell(row=row, column=1, value=line)
        if row == 1:
            cell.font = Font(bold=True, size=13)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="equipment_import_template.xlsx"'
    wb.save(response)
    return response


def _locate_header(ws):
    """Find the header row and map column index -> field name.

    Scans the first :data:`MAX_HEADER_SCAN` rows so files exported by
    :func:`export_equipment_to_excel` (title in A1, headers on row 2) work
    without editing.
    """
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN, values_only=True), start=1
    ):
        mapping = {}
        for col_idx, value in enumerate(row or (), start=1):
            field = HEADER_ALIASES.get(_normalize_header(value))
            if field and field not in mapping.values():
                mapping[col_idx] = field
        if "serial_number" in mapping.values() and "description" in mapping.values():
            return row_idx, mapping
    return None, {}


def _resolve_description(name, create_missing, created_descriptions):
    description = EquipmentDescription.objects.filter(name__iexact=name).first()
    if description:
        return description, None
    if not create_missing:
        return None, f"Equipment description '{name}' does not exist."
    description = EquipmentDescription.objects.create(name=name)
    created_descriptions.append(description)
    return description, None


def _resolve_manufacturer(name, create_missing, created_manufacturers):
    manufacturer = Manufacturer.objects.filter(name__iexact=name).first()
    if manufacturer:
        return manufacturer, None
    if not create_missing:
        return None, f"Manufacturer '{name}' does not exist."
    manufacturer = Manufacturer.objects.create(name=name.title())
    created_manufacturers.append(manufacturer)
    return manufacturer, None


def _process_rows(ws, header_row, mapping, workshop, create_missing, user):
    """Validate and write every data row.

    Returns ``(report, saved)``: ``saved`` lists every row written — auto-created
    descriptions and manufacturers first, then equipment — for the HQ push.
    """
    departments = {
        dept.name.strip().lower(): dept
        for dept in Department.objects.filter(workshop=workshop, active_status=True)
    }

    created_descriptions = []
    created_manufacturers = []
    saved_equipment = []
    rows_report = []
    counts = {"create": 0, "reactivate": 0, "error": 0}
    seen_serials = {}
    total_rows = 0
    truncated = False

    for excel_row, raw in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        values = {field: _cell_text((raw or ())[idx - 1] if raw and idx <= len(raw) else "")
                  for idx, field in mapping.items()}

        # Blank row, or the "Generated by: ..." footer written by the exporter.
        if not any(values.values()):
            continue
        if values.get("description", "").lower().startswith("generated by"):
            continue

        total_rows += 1
        if total_rows > MAX_ROWS:
            truncated = True
            break

        errors = []
        notes = []

        for field, limit in MAX_LENGTHS.items():
            if len(values.get(field, "")) > limit:
                errors.append(f"{COLUMN_LABELS[field]} exceeds {limit} characters.")

        description_name = values.get("description", "")
        model = values.get("model", "")
        serial = values.get("serial_number", "").upper()
        department_name = values.get("department", "")
        manufacturer_name = values.get("manufacturer", "")
        status_raw = values.get("status", "")

        if not description_name:
            errors.append("Description is required.")
        if not model:
            errors.append("Model is required.")
        if not serial:
            errors.append("Serial Number is required.")
        if not department_name:
            errors.append("Department is required.")

        status = STATUS_ALIASES.get(_normalize_status(status_raw))
        if not status:
            if status_raw:
                errors.append(
                    f"Status '{status_raw}' is not valid. Use: {', '.join(STATUS_CHOICES)}."
                )
            else:
                errors.append("Status is required.")

        department = departments.get(department_name.lower()) if department_name else None
        if department_name and not department:
            errors.append(
                f"Department '{department_name}' does not exist in {workshop.name}."
            )

        if serial and serial in seen_serials:
            errors.append(
                f"Serial Number '{serial}' is duplicated in this file (also on row {seen_serials[serial]})."
            )
        elif serial:
            seen_serials[serial] = excel_row

        existing = Equipment.objects.filter(serial_number__iexact=serial).first() if serial else None
        if existing and existing.active_status and not existing.pending_delete:
            errors.append(
                f"Serial Number '{serial}' already exists in {existing.department.name}."
            )

        if errors:
            counts["error"] += 1
            if len(rows_report) < MAX_REPORT_ROWS:
                rows_report.append({
                    "row": excel_row,
                    "serial": serial,
                    "description": description_name,
                    "department": department_name,
                    "action": "error",
                    "messages": errors,
                })
            continue

        # Each row gets its own savepoint so one failure cannot poison the
        # surrounding transaction (and therefore the rest of the import).
        # Lookups created for this row are only kept if its savepoint commits.
        row_descriptions = []
        row_manufacturers = []
        try:
            with transaction.atomic():
                description, desc_error = _resolve_description(
                    description_name, create_missing, row_descriptions
                )
                if desc_error:
                    raise ValueError(desc_error)

                manufacturer = None
                if manufacturer_name:
                    manufacturer, manu_error = _resolve_manufacturer(
                        manufacturer_name, create_missing, row_manufacturers
                    )
                    if manu_error:
                        raise ValueError(manu_error)

                procurement = {"asset_tag": values.get("asset_tag", "")}

                if existing:
                    # Deactivated / pending-delete serial: reactivate in place,
                    # exactly as add_inventory does for the single-item form.
                    existing.description = description
                    existing.manufacturer = manufacturer
                    existing.model = model
                    existing.serial_number = serial
                    existing.department = department
                    existing.status = status
                    for field, value in procurement.items():
                        if value:
                            setattr(existing, field, value)
                    existing.active_status = True
                    existing.pending_delete = False
                    existing.updated_at = timezone.now()
                    existing.save()
                    equipment = existing
                    action = "reactivate"
                    notes.append("Previously deactivated - reactivated and updated.")
                else:
                    equipment = Equipment(
                        description=description,
                        manufacturer=manufacturer,
                        model=model,
                        serial_number=serial,
                        department=department,
                        status=status,
                        active_status=True,
                        pending_delete=False,
                        **procurement,
                    )
                    equipment.save()
                    action = "create"
        except Exception as exc:  # validation or database failure for this row
            counts["error"] += 1
            message = getattr(exc, "messages", None) or [str(exc)]
            if len(rows_report) < MAX_REPORT_ROWS:
                rows_report.append({
                    "row": excel_row,
                    "serial": serial,
                    "description": description_name,
                    "department": department_name,
                    "action": "error",
                    "messages": list(message),
                })
            logger.warning("Equipment import row %s failed for %s: %s", excel_row, user, exc)
            continue

        created_descriptions.extend(row_descriptions)
        created_manufacturers.extend(row_manufacturers)
        saved_equipment.append(equipment)
        counts[action] += 1
        if len(rows_report) < MAX_REPORT_ROWS:
            rows_report.append({
                "row": excel_row,
                "serial": serial,
                "description": description_name,
                "department": department_name,
                "action": action,
                "messages": notes,
            })

    report = {
        "total": total_rows,
        "counts": counts,
        "rows": rows_report,
        "rows_truncated": counts["error"] + counts["create"] + counts["reactivate"] > len(rows_report),
        "file_truncated": truncated,
        "created_descriptions": sorted({d.name for d in created_descriptions}),
        "created_manufacturers": sorted({m.name for m in created_manufacturers}),
    }
    return report, created_descriptions + created_manufacturers + saved_equipment


@login_required
@require_POST
def upload_equipment_excel(request):
    """Preview (default) or commit a bulk equipment upload.

    POST fields:
        ``file``            - the .xlsx workbook (required)
        ``commit``          - "true" to write; anything else previews
        ``create_missing``  - "true" to auto-create descriptions/manufacturers
    """
    profile, workshop, error = _import_scope(request)
    if error:
        return JsonResponse({"success": False, "error": error}, status=403)

    commit = request.POST.get("commit", "").lower() == "true"
    create_missing = request.POST.get("create_missing", "true").lower() == "true"

    # Checked before parsing so an offline user is told straight away. A commit
    # re-checks live instead of trusting the cached answer.
    if not hq_link.is_hq_online(force=commit):
        return JsonResponse({"success": False, "error": HQ_OFFLINE_MESSAGE}, status=503)

    upload = request.FILES.get("file")
    if not upload:
        return JsonResponse({"success": False, "error": "No file was uploaded."}, status=400)

    if not upload.name.lower().endswith((".xlsx", ".xlsm")):
        return JsonResponse(
            {"success": False, "error": "Unsupported file type. Upload an .xlsx workbook."},
            status=400,
        )

    if upload.size > MAX_FILE_SIZE:
        return JsonResponse(
            {"success": False,
             "error": f"File is too large. Maximum size is {MAX_FILE_SIZE // (1024 * 1024)} MB."},
            status=400,
        )

    workbook = None
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Equipment import could not open workbook for %s: %s", request.user, exc)
        return JsonResponse(
            {"success": False, "error": "Could not read the workbook. Make sure it is a valid Excel file."},
            status=400,
        )

    try:
        ws = workbook["Equipment"] if "Equipment" in workbook.sheetnames else workbook.worksheets[0]
        header_row, mapping = _locate_header(ws)

        if not header_row:
            return JsonResponse({
                "success": False,
                "error": (
                    "Could not find the header row. The sheet must contain columns named: "
                    + ", ".join(COLUMN_LABELS[c] for c in COLUMNS)
                    + ". Download the template to get the expected layout."
                ),
            }, status=400)

        missing = [
            COLUMN_LABELS[field]
            for field in ("description", "model", "serial_number", "department", "status")
            if field not in mapping.values()
        ]
        if missing:
            return JsonResponse({
                "success": False,
                "error": "Missing required column(s): " + ", ".join(missing) + ".",
            }, status=400)

        report = None
        try:
            # These rows are pushed to HQ explicitly below, so keep the
            # per-request background push from sending them a second time.
            with hq_link.suppressed(), transaction.atomic():
                report, saved = _process_rows(
                    ws, header_row, mapping, workshop, create_missing, request.user
                )
                if not commit:
                    raise _DryRunRollback()
                # Still inside the transaction: unless HQ accepts every row,
                # the local write is rolled back as well.
                pushed, push_error = hq_link.push_instances(saved)
                if not pushed:
                    raise _HQRejected(push_error)
        except _DryRunRollback:
            pass
        except _HQRejected as exc:
            logger.warning("Equipment import by %s was not accepted by HQ: %s", request.user, exc)
            return JsonResponse(
                {"success": False, "error": f"HQ did not accept the upload, so nothing was saved. {exc}"},
                status=502,
            )
    except Exception as exc:
        logger.error("Equipment import failed for %s: %s", request.user, exc, exc_info=True)
        return JsonResponse(
            {"success": False, "error": "An error occurred while processing the file."},
            status=500,
        )
    finally:
        if workbook is not None:
            workbook.close()

    if report["total"] == 0:
        return JsonResponse(
            {"success": False, "error": "No data rows were found in the workbook."},
            status=400,
        )

    if commit:
        logger.info(
            "Equipment import by %s into %s: %s created, %s reactivated, %s failed.",
            request.user.username, workshop.name,
            report["counts"]["create"], report["counts"]["reactivate"], report["counts"]["error"],
        )

    return JsonResponse({
        "success": True,
        "committed": commit,
        "workshop": workshop.name,
        "filename": upload.name,
        **report,
    })
