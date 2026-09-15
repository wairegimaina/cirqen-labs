"""Tests for the bulk equipment Excel import.

Covers the two endpoints end to end against the in-memory SQLite DB (see
``Equiper.test_settings``): the template download, and the upload's
preview/commit contract. The preview must never write, and the commit must
produce exactly what the preview promised.
"""
import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from openpyxl import Workbook, load_workbook

from Inventory.models import Department, Equipment, EquipmentDescription, Manufacturer
from workshop.models import Workshop
from users.models import UserProfile

User = get_user_model()

HEADERS = ["Description", "Manufacturer", "Model", "Serial Number", "Department", "Status"]


def build_workbook(rows, headers=HEADERS, title_row=None):
    """Return an in-memory .xlsx containing ``rows`` under ``headers``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Equipment"
    if title_row:
        ws.append([title_row])
        ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def as_upload(buffer, name="equipment.xlsx"):
    return SimpleUploadedFile(
        name,
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class ImportTestBase(TestCase):
    def setUp(self):
        self.workshop = Workshop.objects.create(name="Radiology")
        self.other_workshop = Workshop.objects.create(name="Dental")
        self.department = Department.objects.create(name="CT", workshop=self.workshop)
        self.foreign_department = Department.objects.create(
            name="Ortho", workshop=self.other_workshop
        )
        self.description = EquipmentDescription.objects.create(name="Scanner")
        self.manufacturer = Manufacturer.objects.create(name="Acme")

        self.tech = self._make_user("tech", "Tech", workshop=self.workshop, level="Engineer")

        # HQ is online and accepts everything unless a test says otherwise.
        online = mock.patch("core.hq_link.is_hq_online", return_value=True)
        self.is_hq_online = online.start()
        self.addCleanup(online.stop)
        push = mock.patch("core.hq_link.push_instances", return_value=(True, ""))
        self.push_instances = push.start()
        self.addCleanup(push.stop)

    def _make_user(self, username, role, *, workshop=None, department=None, level=None):
        user = User.objects.create_user(username=username, password="pw12345!")
        UserProfile.objects.update_or_create(
            user=user,
            defaults={"role": role, "workshop": workshop, "department": department, "level": level},
        )
        return user

    def upload(self, buffer, *, commit=False, create_missing=True, name="equipment.xlsx"):
        return self.client.post(
            reverse("upload_equipment_excel"),
            {
                "file": as_upload(buffer, name),
                "commit": "true" if commit else "false",
                "create_missing": "true" if create_missing else "false",
            },
        )


class HQConnectionTests(ImportTestBase):
    """Uploads only run while HQ is online, and only commit if HQ accepts them."""

    ROW = ["Monitor", "Globex", "M1", "SN1", "CT", "Working"]

    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def test_upload_refused_while_hq_offline(self):
        self.is_hq_online.return_value = False
        for commit in (False, True):
            resp = self.upload(build_workbook([self.ROW]), commit=commit)
            self.assertEqual(resp.status_code, 503)
            self.assertIn("HQ", resp.json()["error"])
        self.assertEqual(Equipment.objects.count(), 0)
        self.push_instances.assert_not_called()

    def test_commit_checks_hq_live(self):
        self.upload(build_workbook([self.ROW]), commit=True)
        self.is_hq_online.assert_called_once_with(force=True)

    def test_preview_does_not_push(self):
        resp = self.upload(build_workbook([self.ROW]))
        self.assertEqual(resp.status_code, 200)
        self.push_instances.assert_not_called()

    def test_commit_pushes_saved_rows_and_new_lookups(self):
        resp = self.upload(build_workbook([self.ROW]), commit=True)

        self.assertEqual(resp.status_code, 200)
        pushed = self.push_instances.call_args.args[0]
        self.assertEqual(
            [type(obj).__name__ for obj in pushed],
            ["EquipmentDescription", "Manufacturer", "Equipment"],
        )
        self.assertEqual(pushed[2].serial_number, "SN1")

    def test_commit_rolled_back_when_hq_rejects_it(self):
        self.push_instances.return_value = (False, "HQ could not apply 1 of 3 records.")
        resp = self.upload(build_workbook([self.ROW]), commit=True)

        self.assertEqual(resp.status_code, 502)
        self.assertIn("nothing was saved", resp.json()["error"])
        self.assertEqual(Equipment.objects.count(), 0)
        self.assertFalse(EquipmentDescription.objects.filter(name="Monitor").exists())
        self.assertFalse(Manufacturer.objects.filter(name="Globex").exists())


class TemplateDownloadTests(ImportTestBase):
    def test_tech_gets_populated_template(self):
        self.client.force_login(self.tech)
        resp = self.client.get(reverse("download_equipment_import_template"))

        self.assertEqual(resp.status_code, 200)
        self.assertIn("equipment_import_template.xlsx", resp["Content-Disposition"])

        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Equipment", "Reference", "Instructions"])
        self.assertEqual([c.value for c in wb["Equipment"][1]], HEADERS)

    def test_reference_sheet_is_scoped_to_own_workshop(self):
        self.client.force_login(self.tech)
        resp = self.client.get(reverse("download_equipment_import_template"))

        wb = load_workbook(io.BytesIO(resp.content))
        departments = [c.value for c in wb["Reference"]["A"] if c.value][1:]
        self.assertEqual(departments, ["CT"])
        self.assertNotIn("Ortho", departments)

    def test_non_tech_is_forbidden(self):
        hod = self._make_user("hod", "HOD")
        self.client.force_login(hod)
        resp = self.client.get(reverse("download_equipment_import_template"))
        self.assertEqual(resp.status_code, 403)


class UploadPermissionTests(ImportTestBase):
    def test_non_tech_cannot_upload(self):
        hod = self._make_user("hod", "HOD")
        self.client.force_login(hod)
        resp = self.upload(build_workbook([["Scanner", "Acme", "M1", "SN1", "CT", "Working"]]),
                           commit=True)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Equipment.objects.count(), 0)

    def test_tech_without_workshop_cannot_upload(self):
        # UserProfile.clean() forbids a workshop-less Tech, so build a valid one
        # and strip the workshop via the queryset to reach the state the view
        # guards against (e.g. a workshop deleted out from under the profile).
        stray = self._make_user("stray", "Tech", workshop=self.workshop, level="Engineer")
        UserProfile.objects.filter(user=stray).update(workshop=None)
        self.client.force_login(stray)
        resp = self.upload(build_workbook([["Scanner", "Acme", "M1", "SN1", "CT", "Working"]]),
                           commit=True)
        self.assertEqual(resp.status_code, 403)

    def test_anonymous_is_redirected_to_login(self):
        resp = self.upload(build_workbook([["Scanner", "Acme", "M1", "SN1", "CT", "Working"]]))
        self.assertIn(resp.status_code, (302, 301))


class UploadFileValidationTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def test_missing_file_rejected(self):
        resp = self.client.post(reverse("upload_equipment_excel"), {"commit": "false"})
        self.assertEqual(resp.status_code, 400)
        self.assertFalse(resp.json()["success"])

    def test_wrong_extension_rejected(self):
        bad = SimpleUploadedFile("data.csv", b"a,b,c", content_type="text/csv")
        resp = self.client.post(reverse("upload_equipment_excel"), {"file": bad})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Unsupported file type", resp.json()["error"])

    def test_corrupt_workbook_rejected(self):
        bad = SimpleUploadedFile("data.xlsx", b"not really a workbook")
        resp = self.client.post(reverse("upload_equipment_excel"), {"file": bad})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Could not read the workbook", resp.json()["error"])

    def test_missing_required_column_rejected(self):
        buffer = build_workbook(
            [["Scanner", "Acme", "M1", "SN1"]],
            headers=["Description", "Manufacturer", "Model", "Serial Number"],
        )
        resp = self.upload(buffer)
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Missing required column", resp.json()["error"])

    def test_unrecognisable_sheet_rejected(self):
        buffer = build_workbook([["a", "b"]], headers=["Foo", "Bar"])
        resp = self.upload(buffer)
        self.assertEqual(resp.status_code, 400)
        self.assertIn("header row", resp.json()["error"])

    def test_empty_sheet_rejected(self):
        resp = self.upload(build_workbook([]))
        self.assertEqual(resp.status_code, 400)
        self.assertIn("No data rows", resp.json()["error"])


class UploadPreviewTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def test_preview_reports_but_saves_nothing(self):
        buffer = build_workbook([
            ["Scanner", "Acme", "M1", "SN-1", "CT", "Working"],
            ["Scanner", "Acme", "M2", "SN-2", "CT", "Under repair"],
        ])
        resp = self.upload(buffer, commit=False)
        data = resp.json()

        self.assertTrue(data["success"])
        self.assertFalse(data["committed"])
        self.assertEqual(data["total"], 2)
        self.assertEqual(data["counts"]["create"], 2)
        self.assertEqual(data["counts"]["error"], 0)
        self.assertEqual(Equipment.objects.count(), 0)

    def test_preview_does_not_persist_auto_created_lookups(self):
        buffer = build_workbook([["Ventilator", "Draeger", "V500", "SN-9", "CT", "Working"]])
        data = self.upload(buffer, commit=False).json()

        self.assertEqual(data["created_descriptions"], ["Ventilator"])
        self.assertEqual(data["created_manufacturers"], ["Draeger"])
        self.assertFalse(EquipmentDescription.objects.filter(name="Ventilator").exists())
        self.assertFalse(Manufacturer.objects.filter(name="Draeger").exists())

    def test_preview_matches_commit(self):
        rows = [
            ["Scanner", "Acme", "M1", "SN-1", "CT", "Working"],
            ["Scanner", "Acme", "M2", "", "CT", "Working"],  # missing serial
        ]
        preview = self.upload(build_workbook(rows), commit=False).json()
        committed = self.upload(build_workbook(rows), commit=True).json()

        self.assertEqual(preview["counts"], committed["counts"])
        self.assertEqual(Equipment.objects.count(), preview["counts"]["create"])


class UploadCommitTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def test_rows_are_created_with_normalized_serial(self):
        buffer = build_workbook([["Scanner", "Acme", "M1", "sn-lower", "CT", "working"]])
        data = self.upload(buffer, commit=True).json()

        self.assertTrue(data["committed"])
        self.assertEqual(data["counts"]["create"], 1)

        equipment = Equipment.objects.get()
        self.assertEqual(equipment.serial_number, "SN-LOWER")
        self.assertEqual(equipment.status, "Working")
        self.assertEqual(equipment.department, self.department)
        self.assertEqual(equipment.workshop, self.workshop)
        self.assertTrue(equipment.active_status)

    def test_missing_lookups_created_when_opted_in(self):
        buffer = build_workbook([["Ventilator", "Draeger", "V500", "SN-9", "CT", "Working"]])
        self.upload(buffer, commit=True)

        self.assertTrue(EquipmentDescription.objects.filter(name="Ventilator").exists())
        self.assertTrue(Manufacturer.objects.filter(name="Draeger").exists())
        self.assertEqual(Equipment.objects.get().description.name, "Ventilator")

    def test_missing_lookups_are_errors_when_opted_out(self):
        buffer = build_workbook([["Ventilator", "Draeger", "V500", "SN-9", "CT", "Working"]])
        data = self.upload(buffer, commit=True, create_missing=False).json()

        self.assertEqual(data["counts"]["error"], 1)
        self.assertEqual(Equipment.objects.count(), 0)
        self.assertFalse(EquipmentDescription.objects.filter(name="Ventilator").exists())

    def test_existing_lookups_matched_case_insensitively(self):
        buffer = build_workbook([["scanner", "acme", "M1", "SN-1", "ct", "Working"]])
        self.upload(buffer, commit=True)

        equipment = Equipment.objects.get()
        self.assertEqual(equipment.description, self.description)
        self.assertEqual(equipment.manufacturer, self.manufacturer)
        self.assertEqual(equipment.department, self.department)

    def test_optional_manufacturer_may_be_blank(self):
        buffer = build_workbook([["Scanner", "", "M1", "SN-1", "CT", "Working"]])
        data = self.upload(buffer, commit=True).json()

        self.assertEqual(data["counts"]["create"], 1)
        self.assertIsNone(Equipment.objects.get().manufacturer)

    def test_valid_rows_import_alongside_invalid_ones(self):
        buffer = build_workbook([
            ["Scanner", "Acme", "M1", "SN-GOOD", "CT", "Working"],
            ["Scanner", "Acme", "M2", "SN-BAD", "Nowhere", "Working"],
        ])
        data = self.upload(buffer, commit=True).json()

        self.assertEqual(data["counts"]["create"], 1)
        self.assertEqual(data["counts"]["error"], 1)
        self.assertEqual(Equipment.objects.count(), 1)
        self.assertEqual(Equipment.objects.get().serial_number, "SN-GOOD")

    def test_blank_rows_are_ignored(self):
        buffer = build_workbook([
            ["Scanner", "Acme", "M1", "SN-1", "CT", "Working"],
            [None, None, None, None, None, None],
            ["", "", "", "", "", ""],
        ])
        data = self.upload(buffer, commit=True).json()
        self.assertEqual(data["total"], 1)
        self.assertEqual(data["counts"]["create"], 1)


class UploadRowValidationTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def _error_messages(self, row):
        data = self.upload(build_workbook([row]), commit=True).json()
        self.assertEqual(data["counts"]["error"], 1)
        self.assertEqual(Equipment.objects.count(), 0)
        return " ".join(data["rows"][0]["messages"])

    def test_missing_required_fields_reported(self):
        messages = self._error_messages(["", "Acme", "", "", "", ""])
        self.assertIn("Description is required", messages)
        self.assertIn("Model is required", messages)
        self.assertIn("Serial Number is required", messages)
        self.assertIn("Department is required", messages)
        self.assertIn("Status is required", messages)

    def test_invalid_status_reported(self):
        messages = self._error_messages(["Scanner", "Acme", "M1", "SN-1", "CT", "banana"])
        self.assertIn("not valid", messages)

    def test_department_outside_workshop_rejected(self):
        messages = self._error_messages(["Scanner", "Acme", "M1", "SN-1", "Ortho", "Working"])
        self.assertIn("does not exist in Radiology", messages)

    def test_overlong_value_rejected(self):
        messages = self._error_messages(["Scanner", "Acme", "M" * 101, "SN-1", "CT", "Working"])
        self.assertIn("exceeds 100 characters", messages)

    def test_duplicate_serial_within_file_rejected(self):
        buffer = build_workbook([
            ["Scanner", "Acme", "M1", "SN-DUP", "CT", "Working"],
            ["Scanner", "Acme", "M2", "sn-dup", "CT", "Working"],
        ])
        data = self.upload(buffer, commit=True).json()

        self.assertEqual(data["counts"]["create"], 1)
        self.assertEqual(data["counts"]["error"], 1)
        self.assertIn("duplicated in this file", " ".join(data["rows"][1]["messages"]))
        self.assertEqual(Equipment.objects.count(), 1)

    def test_serial_already_active_rejected(self):
        Equipment.objects.create(
            description=self.description, manufacturer=self.manufacturer, model="M0",
            serial_number="SN-EXISTS", department=self.department, status="Working",
        )
        messages = " ".join(
            self.upload(
                build_workbook([["Scanner", "Acme", "M1", "SN-EXISTS", "CT", "Working"]]),
                commit=True,
            ).json()["rows"][0]["messages"]
        )
        self.assertIn("already exists", messages)
        self.assertEqual(Equipment.objects.count(), 1)


class UploadReactivationTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)
        self.retired = Equipment.objects.create(
            description=self.description, manufacturer=self.manufacturer, model="OLD",
            serial_number="SN-OLD", department=self.department, status="Not working",
        )
        Equipment.objects.filter(pk=self.retired.pk).update(
            active_status=False, pending_delete=True
        )

    def test_deactivated_serial_is_reactivated_not_duplicated(self):
        buffer = build_workbook([["Scanner", "Acme", "NEW", "SN-OLD", "CT", "Working"]])
        data = self.upload(buffer, commit=True).json()

        self.assertEqual(data["counts"]["reactivate"], 1)
        self.assertEqual(data["counts"]["create"], 0)
        self.assertEqual(Equipment.objects.count(), 1)

        self.retired.refresh_from_db()
        self.assertTrue(self.retired.active_status)
        self.assertFalse(self.retired.pending_delete)
        self.assertEqual(self.retired.model, "NEW")
        self.assertEqual(self.retired.status, "Working")


class UploadFormatToleranceTests(ImportTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.tech)

    def test_exported_file_layout_round_trips(self):
        """A file shaped like export_equipment_to_excel output is accepted."""
        buffer = build_workbook(
            [
                ["Scanner", "Acme", "M1", "SN-1", "CT", "Working"],
                ["Generated by: someone", None, None, None, None, None],
            ],
            headers=["Description", "Manufacturer", "Model", "Serial No.", "Department", "Status"],
            title_row="Radiology Equipment Inventory",
        )
        data = self.upload(buffer, commit=True).json()

        self.assertEqual(data["total"], 1, "the 'Generated by' footer must be ignored")
        self.assertEqual(data["counts"]["create"], 1)

    def test_numeric_serial_is_not_stringified_as_float(self):
        buffer = build_workbook([["Scanner", "Acme", "M1", 12345, "CT", "Working"]])
        self.upload(buffer, commit=True)
        self.assertEqual(Equipment.objects.get().serial_number, "12345")

    def test_reported_row_numbers_match_the_spreadsheet(self):
        buffer = build_workbook([
            ["Scanner", "Acme", "M1", "SN-1", "CT", "Working"],
            ["Scanner", "Acme", "M2", "SN-2", "Nowhere", "Working"],
        ])
        data = self.upload(buffer, commit=False).json()

        # Headers occupy row 1, so data starts at row 2.
        self.assertEqual([row["row"] for row in data["rows"]], [2, 3])
