"""Shared plumbing for Excel bulk imports of reference data.

Used by the workshop, department, standard and parameter uploads. Each upload
view checks its own permissions, then hands :func:`run_import` a
``process(workbook, report)`` callable that validates the rows and writes the
valid ones. Everything around that follows the same contract as the equipment
import (:mod:`Inventory.views.imports`):

* uploads are refused while HQ is unreachable;
* a preview performs the real writes inside a transaction and rolls them back,
  so what the user is shown is exactly what a commit does;
* a commit is pushed to HQ inside that transaction and rolled back unless HQ
  accepts every row.

Processors write each row inside its own savepoint (``transaction.atomic()``)
so one failing row cannot poison the rest of the import, and must append the
instances they write to ``report.saved`` in foreign-key order (parents first).
"""
import datetime
import logging
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.http import HttpResponse, JsonResponse

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core import hq_link

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 5000                   # per sheet; later rows are not read
MAX_HEADER_SCAN = 15              # rows to search for the header row
MAX_REPORT_ROWS = 500             # per-row detail returned to the browser

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HQ_OFFLINE_MESSAGE = (
    "Bulk upload needs a connection to the HQ server, and HQ is not reachable "
    "right now. Try again once this machine is back online."
)

# Typed (text) dates are read day-first, as they are written locally.
DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%d %b %Y", "%d %B %Y")


class ImportFileError(Exception):
    """A problem with the file as a whole (a missing sheet or header row)."""


class _DryRunRollback(Exception):
    """Raised to unwind the transaction after a preview run."""


class _HQRejected(Exception):
    """Raised to unwind a commit that HQ did not fully accept."""


# ── Cell parsing ─────────────────────────────────────────────────────────────


def normalize(value):
    """``"Serial No."`` -> ``"serialno"``."""
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def cell_text(value):
    """Coerce a cell to trimmed text, without turning whole numbers into floats."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_date(value):
    """Return a ``date`` for a date cell or day-first text; ``None`` when blank.

    Raises ``ValueError`` with a user-facing message.
    """
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = cell_text(value)
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a date. Use a date cell or DD/MM/YYYY.")


def parse_decimal(value, *, max_digits, decimal_places):
    """Parse a cell for a ``DecimalField(max_digits, decimal_places)``.

    Raises ``ValueError`` with a user-facing message.
    """
    if isinstance(value, float):
        # 15 significant digits drops binary noise such as 0.30000000000000004.
        text = format(value, ".15g")
    else:
        text = cell_text(value).lstrip("±").strip()
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"'{text}' is not a number.") from None
    if not number.is_finite():
        raise ValueError(f"'{text}' is not a number.")
    if abs(number) >= Decimal(10) ** (max_digits - decimal_places):
        raise ValueError(f"'{text}' is too large.")
    if number != number.quantize(Decimal(1).scaleb(-decimal_places)):
        raise ValueError(f"'{text}' has more than {decimal_places} decimal places.")
    return number


def error_messages(exc):
    """User-facing messages for an exception raised while saving a row."""
    return list(getattr(exc, "messages", None) or [str(exc)])


# ── Reading sheets ───────────────────────────────────────────────────────────


def find_sheet(workbook, title, fallback_first=False):
    """The sheet named ``title`` (ignoring case and spacing), else ``None``."""
    wanted = normalize(title)
    for name in workbook.sheetnames:
        if normalize(name) == wanted:
            return workbook[name]
    return workbook.worksheets[0] if fallback_first else None


def locate_header(ws, aliases, required):
    """Find the header row and map column index -> field name.

    The header is the first of :data:`MAX_HEADER_SCAN` rows that holds a column
    for every field in ``required``, so a title above the table is tolerated.
    Returns ``(row_number, mapping)``, or ``(None, {})``.
    """
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN, values_only=True), start=1
    ):
        mapping = {}
        for col_idx, value in enumerate(row or (), start=1):
            field = aliases.get(normalize(value))
            if field and field not in mapping.values():
                mapping[col_idx] = field
        if all(field in mapping.values() for field in required):
            return row_idx, mapping
    return None, {}


def header_error(labels, sheet=None):
    where = f"The '{sheet}' sheet" if sheet else "The sheet"
    return (
        f"Could not find the header row. {where} must contain columns named: "
        f"{', '.join(labels)}. Download the template to get the expected layout."
    )


def data_rows(ws, header_row, mapping, report, raw_fields=()):
    """Yield ``(excel_row, values)`` for each non-blank row below the header.

    Values are trimmed text, except ``raw_fields``, which keep the cell's own
    value (dates, numbers). Stops after :data:`MAX_ROWS` rows and flags the
    report as truncated.
    """
    count = 0
    for excel_row, raw in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        raw = raw or ()
        cells = {field: raw[idx - 1] if idx <= len(raw) else None for idx, field in mapping.items()}
        if not any(cell_text(value) for value in cells.values()):
            continue
        count += 1
        if count > MAX_ROWS:
            report.file_truncated = True
            return
        yield excel_row, {
            field: value if field in raw_fields else cell_text(value)
            for field, value in cells.items()
        }


class ImportReport:
    """Per-row outcomes for the browser, plus the rows written for the HQ push."""

    ACTIONS = ("create", "skip", "error")

    def __init__(self):
        self.counts = dict.fromkeys(self.ACTIONS, 0)
        self.rows = []
        self.saved = []      # instances written, parents before children
        self.created = {}    # label -> names auto-created alongside the rows
        self.file_truncated = False

    def record(self, action, *, row, item="", messages=(), sheet=""):
        self.counts[action] += 1
        if len(self.rows) < MAX_REPORT_ROWS:
            self.rows.append({
                "sheet": sheet,
                "row": row,
                "action": action,
                "item": item,
                "messages": list(messages),
            })

    def add_created(self, label, name):
        self.created.setdefault(label, set()).add(name)

    def as_dict(self):
        total = sum(self.counts.values())
        return {
            "total": total,
            "counts": dict(self.counts),
            "rows": self.rows,
            "rows_truncated": total > len(self.rows),
            "file_truncated": self.file_truncated,
            "created": {label: sorted(names) for label, names in self.created.items()},
        }


# ── The upload endpoint ──────────────────────────────────────────────────────


def json_error(message, status):
    return JsonResponse({"success": False, "error": message}, status=status)


def run_import(request, *, what, process):
    """Preview (default) or commit an uploaded workbook.

    POST fields: ``file`` (the .xlsx) and ``commit`` ("true" to write). ``what``
    names the import in logs; ``process(workbook, report)`` does the rows.
    """
    commit = request.POST.get("commit", "").lower() == "true"

    # Checked before parsing so an offline user is told straight away. A commit
    # re-checks live instead of trusting the cached answer.
    if not hq_link.is_hq_online(force=commit):
        return json_error(HQ_OFFLINE_MESSAGE, 503)

    upload = request.FILES.get("file")
    if not upload:
        return json_error("No file was uploaded.", 400)
    if not upload.name.lower().endswith((".xlsx", ".xlsm")):
        return json_error("Unsupported file type. Upload an .xlsx workbook.", 400)
    if upload.size > MAX_FILE_SIZE:
        return json_error(
            f"File is too large. Maximum size is {MAX_FILE_SIZE // (1024 * 1024)} MB.", 400
        )

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("%s import could not open workbook for %s: %s", what, request.user, exc)
        return json_error("Could not read the workbook. Make sure it is a valid Excel file.", 400)

    report = ImportReport()
    try:
        # Saved rows are pushed to HQ explicitly below, so keep the per-request
        # background push from sending them a second time.
        with hq_link.suppressed(), transaction.atomic():
            process(workbook, report)
            if not commit:
                raise _DryRunRollback()
            # Still inside the transaction: unless HQ accepts every row, the
            # local write is rolled back as well.
            if report.saved:
                pushed, push_error = hq_link.push_instances(report.saved)
                if not pushed:
                    raise _HQRejected(push_error)
    except _DryRunRollback:
        pass
    except ImportFileError as exc:
        return json_error(str(exc), 400)
    except _HQRejected as exc:
        logger.warning("%s import by %s was not accepted by HQ: %s", what, request.user, exc)
        return json_error(f"HQ did not accept the upload, so nothing was saved. {exc}", 502)
    except Exception as exc:
        logger.error("%s import failed for %s: %s", what, request.user, exc, exc_info=True)
        return json_error("An error occurred while processing the file.", 500)
    finally:
        workbook.close()

    result = report.as_dict()
    if result["total"] == 0:
        return json_error("No data rows were found in the workbook.", 400)

    if commit:
        logger.info("%s import by %s: %s", what, request.user.username, result["counts"])

    return JsonResponse({"success": True, "committed": commit, "filename": upload.name, **result})


# ── Building templates ───────────────────────────────────────────────────────


def write_header(ws, labels, width=26):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"


def write_reference(wb, columns):
    """Add a 'Reference' sheet with one titled column per ``(title, values)``.

    Returns ``{title: range}`` for the non-empty columns, for :func:`add_dropdown`.
    """
    ref = wb.create_sheet("Reference")
    ranges = {}
    for col, (title, values) in enumerate(columns, start=1):
        letter = get_column_letter(col)
        ref.column_dimensions[letter].width = 34
        ref.cell(row=1, column=col, value=title).font = Font(bold=True)
        for offset, value in enumerate(values, start=2):
            ref.cell(row=offset, column=col, value=value)
        if values:
            ranges[title] = f"'Reference'!${letter}$2:${letter}${len(values) + 1}"
    return ranges


def add_dropdown(ws, column, source_range):
    """Offer the values in ``source_range`` as a dropdown on data rows of ``column``."""
    if not source_range:
        return
    letter = get_column_letter(column)
    validation = DataValidation(type="list", formula1=source_range, allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")


def write_instructions(wb, lines):
    """Add an 'Instructions' sheet, so the data sheets hold nothing but data."""
    notes = wb.create_sheet("Instructions")
    notes.column_dimensions["A"].width = 110
    for row, line in enumerate(lines, start=1):
        cell = notes.cell(row=row, column=1, value=line)
        if row == 1:
            cell.font = Font(bold=True, size=13)


def workbook_response(wb, filename):
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
