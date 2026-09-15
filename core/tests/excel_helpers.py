"""Helpers for testing Excel uploads built on :mod:`core.excel_import`."""
import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile

from openpyxl import Workbook, load_workbook

from core.excel_import import XLSX_CONTENT_TYPE
from users.models import UserProfile

User = get_user_model()


def build_workbook(sheets):
    """Return .xlsx bytes for ``{sheet_title: [row, ...]}`` (header row included)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def as_upload(content, name="import.xlsx"):
    return SimpleUploadedFile(name, content, content_type=XLSX_CONTENT_TYPE)


def read_response_workbook(response):
    return load_workbook(io.BytesIO(response.content))


class ExcelImportTestMixin:
    """HQ is online and accepts every push unless a test says otherwise."""

    def setUp(self):
        super().setUp()
        online = mock.patch("core.hq_link.is_hq_online", return_value=True)
        self.is_hq_online = online.start()
        self.addCleanup(online.stop)
        push = mock.patch("core.hq_link.push_instances", return_value=(True, ""))
        self.push_instances = push.start()
        self.addCleanup(push.stop)

    def make_user(self, username, role, *, workshop=None, department=None, level=None):
        user = User.objects.create_user(username=username, password="pw12345!")
        UserProfile.objects.update_or_create(
            user=user,
            defaults={"role": role, "workshop": workshop, "department": department, "level": level},
        )
        return user

    def post_workbook(self, url, sheets, *, commit=False, **fields):
        return self.client.post(url, {
            "file": as_upload(build_workbook(sheets)),
            "commit": "true" if commit else "false",
            **fields,
        })

    def pushed(self):
        """Instances handed to the HQ push, in order."""
        return list(self.push_instances.call_args.args[0])

    @staticmethod
    def rows_by_number(data):
        return {row["row"]: row for row in data["rows"]}
