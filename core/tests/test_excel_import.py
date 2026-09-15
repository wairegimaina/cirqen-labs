"""Unit tests for cell parsing and sheet reading in :mod:`core.excel_import`.

The upload endpoint itself is exercised through the imports built on it (see
``workshop/test_imports.py``, ``Inventory/tests/test_department_imports.py``
and ``CalSoft/test_imports.py``).
"""
import datetime
from decimal import Decimal
from unittest import mock

from django.test import SimpleTestCase

from openpyxl import Workbook

from core import excel_import as xl


class ParseDateTests(SimpleTestCase):
    def test_date_cells(self):
        self.assertEqual(xl.parse_date(datetime.datetime(2026, 3, 5)), datetime.date(2026, 3, 5))
        self.assertEqual(xl.parse_date(datetime.date(2026, 3, 5)), datetime.date(2026, 3, 5))

    def test_text_is_read_day_first(self):
        for text in ("05/03/2026", "5-3-2026", "05.03.2026", "2026-03-05", "5 Mar 2026"):
            with self.subTest(text=text):
                self.assertEqual(xl.parse_date(text), datetime.date(2026, 3, 5))

    def test_blank_is_none(self):
        self.assertIsNone(xl.parse_date(None))
        self.assertIsNone(xl.parse_date("   "))

    def test_unparseable_text_rejected(self):
        with self.assertRaisesMessage(ValueError, "is not a date"):
            xl.parse_date("next week")


class ParseDecimalTests(SimpleTestCase):
    def parse(self, value):
        return xl.parse_decimal(value, max_digits=10, decimal_places=6)

    def test_numbers_and_text(self):
        self.assertEqual(self.parse(0.5), Decimal("0.5"))
        self.assertEqual(self.parse(2), Decimal("2"))
        self.assertEqual(self.parse(" ±0.001 "), Decimal("0.001"))

    def test_float_noise_is_ignored(self):
        self.assertEqual(self.parse(0.1 + 0.2), Decimal("0.3"))

    def test_too_many_decimal_places_rejected(self):
        with self.assertRaisesMessage(ValueError, "more than 6 decimal places"):
            self.parse("0.0000001")

    def test_too_large_rejected(self):
        self.assertEqual(self.parse("9999.5"), Decimal("9999.5"))
        with self.assertRaisesMessage(ValueError, "too large"):
            self.parse("10000")

    def test_non_numbers_rejected(self):
        for value in ("abc", "nan", "inf"):
            with self.subTest(value=value), self.assertRaisesMessage(ValueError, "not a number"):
                self.parse(value)


class SheetReadingTests(SimpleTestCase):
    ALIASES = {"name": "name", "unit": "unit", "units": "unit"}

    def sheet(self, rows):
        ws = Workbook().active
        for row in rows:
            ws.append(row)
        return ws

    def test_header_found_below_a_title(self):
        ws = self.sheet([["Parameter list"], [], ["Name", "Units"], ["Mass", "kg"]])
        self.assertEqual(
            xl.locate_header(ws, self.ALIASES, required=("name", "unit")),
            (3, {1: "name", 2: "unit"}),
        )

    def test_header_missing_a_required_column(self):
        ws = self.sheet([["Name", "Symbol"], ["Mass", "m"]])
        self.assertEqual(xl.locate_header(ws, self.ALIASES, required=("name", "unit")), (None, {}))

    def test_find_sheet_ignores_case_and_spacing(self):
        wb = Workbook()
        wb.active.title = "Notes"
        target = wb.create_sheet("Standard Parameters")
        self.assertIs(xl.find_sheet(wb, "standardparameters"), target)
        self.assertIsNone(xl.find_sheet(wb, "Standards"))
        self.assertIs(xl.find_sheet(wb, "Standards", fallback_first=True), wb["Notes"])

    def test_data_rows_skip_blanks_and_keep_raw_fields(self):
        ws = self.sheet([["Name", "Unit"], [" Mass ", 1.0], [None, None], ["Length", "m"]])
        rows = list(xl.data_rows(ws, 1, {1: "name", 2: "unit"}, xl.ImportReport(), raw_fields=("unit",)))
        self.assertEqual(rows, [(2, {"name": "Mass", "unit": 1.0}), (4, {"name": "Length", "unit": "m"})])

    def test_data_rows_stop_at_the_row_limit(self):
        ws = self.sheet([["Name"], ["a"], ["b"], ["c"]])
        report = xl.ImportReport()
        with mock.patch.object(xl, "MAX_ROWS", 2):
            rows = list(xl.data_rows(ws, 1, {1: "name"}, report))
        self.assertEqual([number for number, _ in rows], [2, 3])
        self.assertTrue(report.file_truncated)


class ImportReportTests(SimpleTestCase):
    def test_counts_and_detail_cap(self):
        report = xl.ImportReport()
        with mock.patch.object(xl, "MAX_REPORT_ROWS", 1):
            report.record("create", row=2, item="a")
            report.record("error", row=3, item="b", messages=["bad"])
        report.add_created("Parameters", "Mass (kg)")
        data = report.as_dict()
        self.assertEqual(data["total"], 2)
        self.assertEqual(data["counts"], {"create": 1, "skip": 0, "error": 1})
        self.assertEqual(len(data["rows"]), 1)
        self.assertTrue(data["rows_truncated"])
        self.assertEqual(data["created"], {"Parameters": ["Mass (kg)"]})
