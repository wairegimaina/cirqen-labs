"""Bulk import of calibration standards and parameters from Excel.

* Parameters: one row per parameter (Name, Symbol, Unit).
* Standards: a 'Standards' sheet with one row per standard, and a
  'Standard Parameters' sheet with one row per parameter a standard provides,
  linked by the standard's serial number. A standard is imported together with
  all of its parameter rows or not at all, so it is never saved half-described.
  Parameters named there that don't exist yet can be created on the fly.

Procedures are deliberately not importable: their parameters, sub-parameters
and set values need the ordering and checks of the procedure form.

Uploads follow the shared preview/commit contract in :mod:`core.excel_import`.
"""
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.views.decorators.http import require_GET, require_POST

from openpyxl import Workbook

from CalSoft.models import Parameter, Standard, StandardParameter
from core import excel_import as xl

IMPORT_ROLES = ("Tech", "HOD")

STANDARDS_SHEET = "Standards"
LINKS_SHEET = "Standard Parameters"


def _max_length(model, field):
    return model._meta.get_field(field).max_length


def _denied(request):
    profile = getattr(request.user, "userprofile", None)
    if profile is None or profile.role not in IMPORT_ROLES:
        return xl.json_error(
            "Only Technologists and HODs can bulk-upload standards and parameters.", 403
        )
    return None


def _check_lengths(values, limits, labels, errors):
    for field, limit in limits.items():
        if len(values.get(field) or "") > limit:
            errors.append(f"{labels[field]} exceeds {limit} characters.")


def _active_parameters():
    return Parameter.objects.filter(pending_delete=False, active_status=True)


def _find_parameter(name, unit):
    """Match by name (any case) and unit (exact case: 'mV' is not 'MV')."""
    return _active_parameters().filter(name__iexact=name, unit=unit).first()


def _write_parameter_reference(wb, extra=()):
    existing = list(_active_parameters().order_by("name", "unit").values_list("name", "unit", "symbol"))
    xl.write_reference(wb, [
        ("Existing Parameter", [name for name, _, _ in existing]),
        ("Unit", [unit for _, unit, _ in existing]),
        ("Symbol", [symbol for _, _, symbol in existing]),
        *extra,
    ])


# ── Parameters ───────────────────────────────────────────────────────────────

PARAMETER_COLUMNS = {"name": "Name", "symbol": "Symbol", "unit": "Unit"}
PARAMETER_REQUIRED = ("name", "unit")
PARAMETER_ALIASES = {
    "name": "name",
    "parameter": "name",
    "parametername": "name",
    "symbol": "symbol",
    "unit": "unit",
    "units": "unit",
    "uom": "unit",
    "unitofmeasure": "unit",
}
PARAMETER_LIMITS = {field: _max_length(Parameter, field) for field in PARAMETER_COLUMNS}


@login_required
@require_GET
def download_parameters_import_template(request):
    denied = _denied(request)
    if denied:
        return denied

    wb = Workbook()
    ws = wb.active
    ws.title = "Parameters"
    xl.write_header(ws, list(PARAMETER_COLUMNS.values()))
    _write_parameter_reference(wb)
    xl.write_instructions(wb, [
        "How to use this template",
        "",
        "1. Enter one parameter per row on the 'Parameters' sheet, starting at row 2.",
        "2. Do not rename, reorder or delete the header row.",
        "3. Name and Unit are required. Symbol is optional.",
        f"   Limits: Name {PARAMETER_LIMITS['name']}, Symbol {PARAMETER_LIMITS['symbol']}, "
        f"Unit {PARAMETER_LIMITS['unit']} characters.",
        "4. A parameter is identified by Name and Unit together, so 'Temperature' in °C and in K are two",
        "   parameters. Units are case-sensitive: mV (millivolt) is not MV (megavolt).",
        "5. Parameters that already exist (see the 'Reference' sheet) are skipped, not changed.",
        f"6. Maximum {xl.MAX_ROWS} rows per upload.",
        "",
        "You will always see a preview of what will be imported before anything is saved.",
    ])
    return xl.workbook_response(wb, "parameters_import_template.xlsx")


@login_required
@require_POST
def upload_parameters_excel(request):
    denied = _denied(request)
    if denied:
        return denied
    return xl.run_import(
        request,
        what="Parameter",
        process=lambda workbook, report: _process_parameters(workbook, report, request.user),
    )


def _process_parameters(workbook, report, user):
    ws = xl.find_sheet(workbook, "Parameters", fallback_first=True)
    header_row, mapping = xl.locate_header(ws, PARAMETER_ALIASES, required=PARAMETER_REQUIRED)
    if not header_row:
        raise xl.ImportFileError(xl.header_error([PARAMETER_COLUMNS[f] for f in PARAMETER_REQUIRED]))

    seen = {}
    for excel_row, values in xl.data_rows(ws, header_row, mapping, report):
        name, unit, symbol = values["name"], values["unit"], values.get("symbol", "")
        item = f"{name} ({unit})" if unit else name
        errors = []

        if not name:
            errors.append("Name is required.")
        if not unit:
            errors.append("Unit is required.")
        _check_lengths(values, PARAMETER_LIMITS, PARAMETER_COLUMNS, errors)

        key = (name.lower(), unit)
        if name and unit and key in seen:
            errors.append(f"'{item}' is duplicated in this file (also on row {seen[key]}).")
        elif name and unit:
            seen[key] = excel_row

        if errors:
            report.record("error", row=excel_row, item=item, messages=errors)
            continue

        if _find_parameter(name, unit):
            report.record("skip", row=excel_row, item=item, messages=["Already exists - not changed."])
            continue

        try:
            with transaction.atomic():
                parameter = Parameter.objects.create(name=name, symbol=symbol, unit=unit, created_by=user)
        except Exception as exc:
            report.record("error", row=excel_row, item=item, messages=xl.error_messages(exc))
            continue

        report.saved.append(parameter)
        report.record("create", row=excel_row, item=item)


# ── Standards ────────────────────────────────────────────────────────────────

STANDARD_COLUMNS = {
    "name": "Name",
    "model_number": "Model Number",
    "serial_number": "Serial Number",
    "manufacturer": "Manufacturer",
    "certificate_number": "Certificate Number",
    "calibration_date": "Calibration Date",
    "calibration_due_date": "Due Date",
    "calibration_agency": "Calibration Agency",
}
STANDARD_REQUIRED = (
    "name", "model_number", "serial_number", "manufacturer", "calibration_date", "calibration_due_date",
)
STANDARD_DATES = ("calibration_date", "calibration_due_date")
STANDARD_ALIASES = {
    "name": "name",
    "standardname": "name",
    "model": "model_number",
    "modelnumber": "model_number",
    "modelno": "model_number",
    "serialnumber": "serial_number",
    "serialno": "serial_number",
    "serial": "serial_number",
    "sn": "serial_number",
    "manufacturer": "manufacturer",
    "make": "manufacturer",
    "brand": "manufacturer",
    "certificatenumber": "certificate_number",
    "certificateno": "certificate_number",
    "certno": "certificate_number",
    "certificate": "certificate_number",
    "calibrationdate": "calibration_date",
    "caldate": "calibration_date",
    "lastcalibration": "calibration_date",
    "lastcalibrated": "calibration_date",
    "duedate": "calibration_due_date",
    "calibrationduedate": "calibration_due_date",
    "calduedate": "calibration_due_date",
    "nextcalibration": "calibration_due_date",
    "nextcalibrationdate": "calibration_due_date",
    "calibrationagency": "calibration_agency",
    "agency": "calibration_agency",
    "calibratedby": "calibration_agency",
}
STANDARD_LIMITS = {
    field: _max_length(Standard, field) for field in STANDARD_COLUMNS if field not in STANDARD_DATES
}

LINK_COLUMNS = {
    "serial_number": "Standard Serial Number",
    "parameter": "Parameter",
    "unit": "Unit",
    "symbol": "Symbol",
    "uncertainty": "Uncertainty (k=2)",
}
LINK_REQUIRED = ("serial_number", "parameter", "unit", "uncertainty")
LINK_ALIASES = {
    "standardserialnumber": "serial_number",
    "standardserialno": "serial_number",
    "standardserial": "serial_number",
    "serialnumber": "serial_number",
    "serialno": "serial_number",
    "serial": "serial_number",
    "parameter": "parameter",
    "parametername": "parameter",
    "unit": "unit",
    "units": "unit",
    "uom": "unit",
    "symbol": "symbol",
    "uncertainty": "uncertainty",
    "uncertaintyk2": "uncertainty",
    "expandeduncertainty": "uncertainty",
}
LINK_LIMITS = {
    "serial_number": _max_length(Standard, "serial_number"),
    "parameter": _max_length(Parameter, "name"),
    "unit": _max_length(Parameter, "unit"),
    "symbol": _max_length(Parameter, "symbol"),
}
UNCERTAINTY_FIELD = StandardParameter._meta.get_field("uncertainty")


@login_required
@require_GET
def download_standards_import_template(request):
    denied = _denied(request)
    if denied:
        return denied

    wb = Workbook()
    ws = wb.active
    ws.title = STANDARDS_SHEET
    xl.write_header(ws, list(STANDARD_COLUMNS.values()))
    xl.write_header(wb.create_sheet(LINKS_SHEET), list(LINK_COLUMNS.values()))

    serials = Standard.objects.filter(pending_delete=False).order_by("serial_number")
    _write_parameter_reference(wb, extra=[
        ("Existing Standard Serial Numbers", list(serials.values_list("serial_number", flat=True))),
    ])

    xl.write_instructions(wb, [
        "How to use this template",
        "",
        "The workbook has two sheets that work together:",
        f"  '{STANDARDS_SHEET}' - one row per reference standard.",
        f"  '{LINKS_SHEET}' - one row for each parameter a standard measures, with its uncertainty.",
        "  The two are linked by the standard's Serial Number.",
        "",
        "1. Do not rename the sheets, and do not rename, reorder or delete the header rows.",
        "2. Name, Model Number, Serial Number, Manufacturer, Calibration Date and Due Date are required.",
        "   Certificate Number and Calibration Agency are optional; a certificate number is generated if blank.",
        "3. Enter dates as Excel dates, or as text in DD/MM/YYYY format. Due Date must be after Calibration Date.",
        "4. Serial Number must be unique. Standards that already exist are skipped, not changed.",
        f"5. Every standard needs at least one row on '{LINKS_SHEET}'. If any of a standard's parameter rows",
        "   has an error, that standard is not imported at all.",
        "   A standard may list each parameter name only once, even in different units, because its",
        "   uncertainty is looked up by parameter name when procedures are set up.",
        "6. Uncertainty (k=2) must be a positive number with at most 6 decimal places.",
        "7. Parameters are matched by Name and Unit. Units are case-sensitive: mV is not MV.",
        "   Parameters that do not exist yet can be created automatically during upload (see 'Reference').",
        f"8. Maximum {xl.MAX_ROWS} rows per sheet.",
        "",
        "You will always see a preview of what will be imported before anything is saved.",
    ])
    return xl.workbook_response(wb, "standards_import_template.xlsx")


@login_required
@require_POST
def upload_standards_excel(request):
    """POST fields as :func:`core.excel_import.run_import`, plus ``create_missing``
    ("true" to create parameters that don't exist yet)."""
    denied = _denied(request)
    if denied:
        return denied
    create_missing = request.POST.get("create_missing", "true").lower() == "true"
    return xl.run_import(
        request,
        what="Standard",
        process=lambda workbook, report: _process_standards(
            workbook, report, request.user, create_missing
        ),
    )


def _read_links(ws, header_row, mapping, report):
    """Validate the 'Standard Parameters' rows, grouped by lower-cased serial number."""
    links = {}
    seen = {}
    for excel_row, values in xl.data_rows(ws, header_row, mapping, report, raw_fields=("uncertainty",)):
        serial, name, unit = values["serial_number"], values["parameter"], values["unit"]
        errors = []

        if not serial:
            errors.append("Standard Serial Number is required.")
        if not name:
            errors.append("Parameter is required.")
        if not unit:
            errors.append("Unit is required.")
        _check_lengths(values, LINK_LIMITS, LINK_COLUMNS, errors)

        uncertainty = None
        if not xl.cell_text(values["uncertainty"]):
            errors.append("Uncertainty is required.")
        else:
            try:
                uncertainty = xl.parse_decimal(
                    values["uncertainty"],
                    max_digits=UNCERTAINTY_FIELD.max_digits,
                    decimal_places=UNCERTAINTY_FIELD.decimal_places,
                )
            except ValueError as exc:
                errors.append(f"Uncertainty: {exc}")
            else:
                if uncertainty <= 0:
                    errors.append("Uncertainty must be positive.")

        # CalSoft looks a standard's uncertainty up by parameter *name* (the
        # procedure form's auto-fill, CalibrationProcedure.get_parameters_with_standards),
        # so a standard must not list the same name twice, even in two units.
        key = (serial.lower(), name.lower())
        if serial and name and key in seen:
            errors.append(
                f"Parameter '{name}' is listed twice for this standard (also on row {seen[key]}). "
                "A standard can have each parameter name only once, because its uncertainty is looked up by name."
            )
        elif serial and name:
            seen[key] = excel_row

        links.setdefault(serial.lower(), []).append({
            "row": excel_row,
            "serial": serial,
            "parameter": name,
            "unit": unit,
            "symbol": values.get("symbol", ""),
            "uncertainty": uncertainty,
            "errors": errors,
        })
    return links


def _process_standards(workbook, report, user, create_missing):
    standards_ws = xl.find_sheet(workbook, STANDARDS_SHEET, fallback_first=True)
    links_ws = xl.find_sheet(workbook, LINKS_SHEET)
    if links_ws is None:
        raise xl.ImportFileError(
            f"The workbook needs a '{LINKS_SHEET}' sheet listing each standard's parameters. "
            "Download the template to get the expected layout."
        )

    header_row, mapping = xl.locate_header(standards_ws, STANDARD_ALIASES, required=STANDARD_REQUIRED)
    if not header_row:
        raise xl.ImportFileError(xl.header_error(
            [STANDARD_COLUMNS[f] for f in STANDARD_REQUIRED], sheet=STANDARDS_SHEET
        ))
    links_header_row, links_mapping = xl.locate_header(links_ws, LINK_ALIASES, required=LINK_REQUIRED)
    if not links_header_row:
        raise xl.ImportFileError(xl.header_error(
            [LINK_COLUMNS[f] for f in LINK_REQUIRED], sheet=LINKS_SHEET
        ))

    links = _read_links(links_ws, links_header_row, links_mapping, report)

    # Collected per table so the HQ push sends parents before children.
    new_parameters, new_standards, new_links = [], [], []
    seen_serials = {}

    for excel_row, values in xl.data_rows(
        standards_ws, header_row, mapping, report, raw_fields=STANDARD_DATES
    ):
        serial = values["serial_number"]
        item = f"{values['name']} ({serial})" if serial else values["name"]
        errors = []

        for field in STANDARD_REQUIRED:
            if field not in STANDARD_DATES and not values[field]:
                errors.append(f"{STANDARD_COLUMNS[field]} is required.")
        _check_lengths(values, STANDARD_LIMITS, STANDARD_COLUMNS, errors)

        dates = {}
        for field in STANDARD_DATES:
            try:
                dates[field] = xl.parse_date(values[field])
            except ValueError as exc:
                errors.append(f"{STANDARD_COLUMNS[field]}: {exc}")
                continue
            if dates[field] is None:
                errors.append(f"{STANDARD_COLUMNS[field]} is required.")
        calibration_date = dates.get("calibration_date")
        due_date = dates.get("calibration_due_date")
        if calibration_date and due_date and due_date <= calibration_date:
            errors.append("Due Date must be after Calibration Date.")

        rows = []
        duplicate = bool(serial) and serial.lower() in seen_serials
        if duplicate:
            errors.append(
                f"Serial Number '{serial}' is duplicated in this file (also on row {seen_serials[serial.lower()]})."
            )
        elif serial:
            seen_serials[serial.lower()] = excel_row
            rows = links.pop(serial.lower(), [])

        existing = Standard.objects.filter(serial_number__iexact=serial).first() if serial else None
        if existing and existing.pending_delete:
            errors.append(
                f"Standard '{existing.serial_number}' is pending deletion. Wait for it to sync before adding it again."
            )
        elif existing and not errors:
            report.record("skip", sheet=STANDARDS_SHEET, row=excel_row, item=item, messages=[
                "Already exists - not changed. Update its details on the Standards page."
            ])
            continue
        elif existing:
            errors.append(f"Standard '{existing.serial_number}' already exists.")

        if serial and not existing and not duplicate and not rows:
            errors.append(f"No parameters are listed for this standard on the '{LINKS_SHEET}' sheet.")
        for link in rows:
            errors.extend(f"{LINKS_SHEET} row {link['row']}: {message}" for message in link["errors"])

        if errors:
            report.record("error", sheet=STANDARDS_SHEET, row=excel_row, item=item, messages=errors)
            continue

        row_parameters = []
        row_links = []
        notes = []
        try:
            with transaction.atomic():
                standard = Standard(
                    name=values["name"],
                    model_number=values["model_number"],
                    serial_number=serial,
                    manufacturer=values["manufacturer"],
                    certificate_number=values.get("certificate_number", ""),
                    calibration_date=calibration_date,
                    calibration_due_date=due_date,
                    calibration_agency=values.get("calibration_agency", ""),
                    created_by=user,
                )
                standard.save()
                for link in rows:
                    parameter = _find_parameter(link["parameter"], link["unit"])
                    if parameter is None:
                        if not create_missing:
                            raise ValueError(
                                f"{LINKS_SHEET} row {link['row']}: Parameter "
                                f"'{link['parameter']} ({link['unit']})' does not exist."
                            )
                        # Units are case-sensitive (mV/MV), so this is only flagged:
                        # 'mmhg' next to an existing 'mmHg' is most likely a typo.
                        other_units = sorted(set(
                            _active_parameters()
                            .filter(name__iexact=link["parameter"])
                            .values_list("unit", flat=True)
                        ))
                        if other_units:
                            notes.append(
                                f"Check the unit of new parameter '{link['parameter']} ({link['unit']})': "
                                f"'{link['parameter']}' already exists in: {', '.join(other_units)}."
                            )
                        parameter = Parameter.objects.create(
                            name=link["parameter"], symbol=link["symbol"], unit=link["unit"], created_by=user
                        )
                        row_parameters.append(parameter)
                    row_links.append(StandardParameter.objects.create(
                        standard=standard, parameter=parameter, uncertainty=link["uncertainty"]
                    ))
        except Exception as exc:
            report.record("error", sheet=STANDARDS_SHEET, row=excel_row, item=item,
                          messages=xl.error_messages(exc))
            continue

        new_parameters.extend(row_parameters)
        new_standards.append(standard)
        new_links.extend(row_links)
        for parameter in row_parameters:
            report.add_created("Parameters", f"{parameter.name} ({parameter.unit})")
        report.record("create", sheet=STANDARDS_SHEET, row=excel_row, item=item, messages=[
            f"Due {due_date:%d/%m/%Y}",
            "Parameters: " + ", ".join(
                f"{link.parameter.name} ({link.parameter.unit}) ±{link.uncertainty.normalize():f}"
                for link in row_links
            ),
            *notes,
        ])

    # Parameter rows whose serial number matched no row on the Standards sheet.
    for rows in links.values():
        for link in rows:
            messages = list(link["errors"])
            if link["serial"]:
                if Standard.objects.filter(serial_number__iexact=link["serial"]).exists():
                    messages.insert(0, (
                        f"Standard '{link['serial']}' is not on the '{STANDARDS_SHEET}' sheet. "
                        "Add parameters to an existing standard on the Standards page."
                    ))
                else:
                    messages.insert(0, f"Serial Number '{link['serial']}' is not on the '{STANDARDS_SHEET}' sheet.")
            report.record("error", sheet=LINKS_SHEET, row=link["row"],
                          item=f"{link['parameter']} ({link['unit']})", messages=messages)

    report.saved.extend(new_parameters + new_standards + new_links)
