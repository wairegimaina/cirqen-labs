import io
from datetime import datetime
from django.contrib.auth import get_user_model
from django.test import TestCase, Client
from django.urls import reverse
from django.http import HttpResponse
from openpyxl import load_workbook

from calSchedules.models import CalibrationSchedule
from Inventory.models import Department, Equipment, EquipmentDescription
from users.models import UserProfile
from workshop.models import Workshop

User = get_user_model()


class CalibrationViewsTest(TestCase):
    def setUp(self):
        # Users
        self.dept_user = User.objects.create_user(username="deptuser", password="pass")
        self.admin_user = User.objects.create_superuser(username="admin", password="pass")

        # Workshop + Department
        # Calibration is planned by the calibration center (scheduling.planner).
        self.workshop = Workshop.objects.create(name="Main Workshop", category="calibration_center")
        self.department = Department.objects.create(name="Lab A", workshop=self.workshop)

        # calSchedules permissions come from the UserProfile, not is_superuser:
        # an Engineer Incharge gets workshop-wide edit + schedule access.
        UserProfile.objects.update_or_create(
            user=self.admin_user,
            defaults={"role": "Tech", "level": "Engineer Incharge", "workshop": self.workshop},
        )

        # Equipment + Description
        self.description = EquipmentDescription.objects.create(name="Balance Scale")
        self.equipment = Equipment.objects.create(
            department=self.department,
            workshop=self.workshop,
            description=self.description,
            model="M123",
            serial_number="SN001",
            status="Working",
        )

        # Calibration Schedule
        self.schedule = CalibrationSchedule.objects.create(
            equipment=self.equipment,
            workshop=self.workshop,
            scheduled_month=datetime.today().replace(day=1),
            status="pending",
            calibration_period=12,
        )

        self.client = Client()

    def login_as_dept_user(self):
        self.client.login(username="deptuser", password="pass")

    def login_as_admin(self):
        self.client.login(username="admin", password="pass")

    # -----------------------
    # Delete / Bulk Delete
    # -----------------------
    def test_delete_schedule_admin(self):
        self.login_as_admin()
        response = self.client.post(reverse("schedule:delete_calibration_schedule", args=[self.schedule.id]))
        self.assertRedirects(
            response, reverse("schedule:calibration_dashboard"), fetch_redirect_response=False
        )
        # Deletes are soft (pending_delete) so the sync agent can propagate them to HQ.
        self.schedule.refresh_from_db()
        self.assertTrue(self.schedule.pending_delete)

    def test_bulk_delete_schedule_admin(self):
        self.login_as_admin()
        response = self.client.post(reverse("schedule:bulk_delete_calibration_schedules"), {
            "schedule_ids": [self.schedule.id],
        })
        self.assertRedirects(
            response, reverse("schedule:calibration_dashboard"), fetch_redirect_response=False
        )
        # Deletes are soft (pending_delete) so the sync agent can propagate them to HQ.
        self.schedule.refresh_from_db()
        self.assertTrue(self.schedule.pending_delete)

    # -----------------------
    # Push / Bulk Push
    # -----------------------
    def test_bulk_push_schedule_admin(self):
        self.login_as_admin()
        scheduled_before = self.schedule.scheduled_month
        response = self.client.post(reverse("schedule:bulk_push_calibration_schedules"), {
            "schedule_ids": [self.schedule.id],
        })
        self.assertRedirects(
            response, reverse("schedule:calibration_dashboard"), fetch_redirect_response=False
        )
        self.schedule.refresh_from_db()
        self.assertNotEqual(self.schedule.scheduled_month, scheduled_before)
        self.assertEqual(self.schedule.status, "pushed")

    # -----------------------
    # Edit Schedule
    # -----------------------
    def test_edit_schedule_admin(self):
        self.login_as_admin()
        new_month = datetime.today().replace(day=1)
        response = self.client.post(
            reverse("schedule:edit_calibration_schedule", args=[self.schedule.id]),
            {
                "scheduled_month": new_month.strftime("%Y-%m"),
                "status": "pending",
                "calibration_period": 6,
            },
        )
        self.assertRedirects(
            response, reverse("schedule:calibration_dashboard"), fetch_redirect_response=False
        )
        self.schedule.refresh_from_db()
        self.assertEqual(self.schedule.calibration_period, 6)

    # -----------------------
    # Schedule New Equipment
    # -----------------------
    def test_schedule_equipment_admin(self):
        self.login_as_admin()
        new_eq = Equipment.objects.create(
            department=self.department,
            workshop=self.workshop,
            description=self.description,
            model="M124",
            serial_number="SN002",
            status="Working",
        )
        response = self.client.post(reverse("schedule:schedule_calibration_equipment", args=[new_eq.id]))
        self.assertRedirects(
            response, reverse("schedule:calibration_dashboard"), fetch_redirect_response=False
        )
        self.assertTrue(CalibrationSchedule.objects.filter(equipment=new_eq).exists())

    # -----------------------
    # Export to Excel
    # -----------------------
    def test_export_calibration_excel_admin(self):
        self.login_as_admin()
        response = self.client.get(reverse("schedule:export_calibration_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(response, HttpResponse)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # Parse workbook to ensure content
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Description", headers)
        self.assertTrue(ws.max_row >= 2)  # at least one schedule exported
