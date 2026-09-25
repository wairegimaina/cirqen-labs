"""machineReports views — Excel/PDF exports and their ReportLab doc templates."""
from uuid import UUID
from django.shortcuts import redirect, render, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.db.models import Count, Sum, Avg, Q, F
from django.utils import timezone
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.contrib import messages
from datetime import datetime, timedelta
import json
import logging
from openpyxl import Workbook
from collections import defaultdict
from django.conf import settings
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods
from django.shortcuts import get_object_or_404
from django.db.models import Sum, Avg, Count, Q
from Inventory.models import Equipment
from CalSoft.models import CalibrationSession
from ..models import EquipmentStatusReport, MachineRepairHistory, WorkshopEquipmentReport, EquipmentCategory
from Inventory.models import Equipment, Workshop
from jobcard.models import jobcard, SparePartUsed
from CalSoft.models import CalibrationSession, CalibrationSchedule
logger = logging.getLogger(__name__)
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from ..models import EquipmentCategory
from Inventory.models import EquipmentDescription
from django.template.loader import render_to_string
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, BaseDocTemplate, Frame, PageTemplate
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, cm
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from django.http import HttpResponse
from django.utils import timezone
from django.contrib.auth.decorators import login_required
import os
from django.conf import settings
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Q
from ..equipment_pdf_generator import create_equipment_pdf_response
from Inventory.models import Equipment, Workshop
from ..manufacturer_performance_pdf_generator import create_manufacturer_pdf_response
from Inventory.models import Equipment, Workshop

# sibling modules in this package
from .helpers import calculate_manufacturer_performance, get_user_workshop_context
from core.branding import contact_line
from core.eat import fmt_eat


@login_required
def export_equipment_history(request, equipment_id):
    """Export equipment history to Excel with complete cost breakdown"""
    profile, workshops, selected_workshop, is_hod = get_user_workshop_context(request)

    if profile is None:
        messages.error(request, "You do not have a user profile assigned. Please contact the administrator.")
        return redirect('custom_login')

    try:
        equipment = Equipment.objects.select_related('workshop', 'category', 'description', 'manufacturer').filter(active_status=True).get(id=equipment_id)
    except Equipment.DoesNotExist:
        messages.error(request, "Equipment not found.")
        return redirect('equipment_dashboard')

    if not is_hod and equipment.workshop != selected_workshop:
        messages.error(request, "You do not have permission to export this equipment's history.")
        return redirect('equipment_dashboard')

    # Create workbook
    wb = Workbook()

    # Equipment Details Sheet
    ws_details = wb.active
    ws_details.title = "Equipment Details"
    ws_details.append(["Equipment Details"])
    ws_details.append(["Name", equipment.description.name if equipment.description else "Unknown"])
    ws_details.append(["Serial Number", equipment.serial_number or "N/A"])
    ws_details.append(["Category", equipment.category.name if equipment.category else "Uncategorized"])
    ws_details.append(["Workshop", equipment.workshop.name if equipment.workshop else "N/A"])
    ws_details.append(["Status", equipment.status or "Unknown"])
    ws_details.append(["Manufacturer", equipment.manufacturer.name if equipment.manufacturer else "Unknown"])
    ws_details.append(["Model", equipment.model or "N/A"])
    ws_details.append(["Description", str(equipment.description) if equipment.description else "No description"])

    # Repair History Sheet with complete costs
    ws_repairs = wb.create_sheet("Repair History")
    ws_repairs.append([
        "Date", "Work Order ID", "Description", "Action Taken", "Performed By",
        "Time Started", "Time Completed", "Downtime (Hours)",
        "Labor Cost (KSh)", "Parts Cost (KSh)", "Additional Costs (KSh)",
        "Additional Costs Description", "Total Cost (KSh)", "Spare Parts Details"
    ])

    repairs = jobcard.objects.filter(
        equipment=equipment,
        status='Approved',
        action_taken='Repair'
    ).prefetch_related('spare_parts__part').order_by('-date_issued')

    total_repairs = 0
    total_downtime = 0
    total_labor = 0
    total_parts = 0
    total_additional = 0
    total_cost = 0

    for repair in repairs:
        # Build spare parts details string
        spare_parts_details = []
        for sp in repair.spare_parts.all():
            part_name = sp.part.name.name if sp.part and sp.part.name else 'Unknown'
            part_cost = float(sp.get_total_cost())
            spare_parts_details.append(
                f"{part_name} (x{sp.quantity} @ KSh {float(sp.unit_cost):.2f} = KSh {part_cost:.2f})"
            )
        spare_parts_str = "; ".join(spare_parts_details) if spare_parts_details else "No parts used"

        # Calculate downtime
        downtime_hours = 0
        if repair.time_started and repair.time_completed:
            start_time = datetime.combine(repair.date_issued, repair.time_started)
            end_time = datetime.combine(repair.date_issued, repair.time_completed)
            downtime = end_time - start_time
            downtime_hours = round(downtime.total_seconds() / 3600, 2)

        # Get costs
        labor_cost = float(repair.labor_cost)
        parts_cost = float(repair.total_parts_cost)
        additional_costs = float(repair.additional_costs)
        repair_total = float(repair.get_total_cost())

        # Accumulate totals
        total_repairs += 1
        total_downtime += downtime_hours
        total_labor += labor_cost
        total_parts += parts_cost
        total_additional += additional_costs
        total_cost += repair_total

        ws_repairs.append([
            repair.date_issued.strftime('%Y-%m-%d'),
            str(repair.id),
            repair.job_description,
            repair.action_taken,
            repair.performed_by.get_full_name() if repair.performed_by else 'N/A',
            repair.time_started.strftime('%H:%M') if repair.time_started else 'N/A',
            repair.time_completed.strftime('%H:%M') if repair.time_completed else 'N/A',
            downtime_hours,
            labor_cost,
            parts_cost,
            additional_costs,
            repair.additional_costs_description or 'N/A',
            repair_total,
            spare_parts_str
        ])

    # Add summary rows
    ws_repairs.append([])  # Empty row
    ws_repairs.append(["SUMMARY"])
    ws_repairs.append(["Total Repairs:", total_repairs])
    ws_repairs.append(["Total Downtime (hours):", round(total_downtime, 2)])
    ws_repairs.append(["Total Labor Cost (KSh):", round(total_labor, 2)])
    ws_repairs.append(["Total Parts Cost (KSh):", round(total_parts, 2)])
    ws_repairs.append(["Total Additional Costs (KSh):", round(total_additional, 2)])
    ws_repairs.append(["Total Cost (KSh):", round(total_cost, 2)])
    if total_repairs > 0:
        ws_repairs.append(["Average Cost per Repair (KSh):", round(total_cost / total_repairs, 2)])

    # Cost Analysis Sheet
    ws_cost_analysis = wb.create_sheet("Cost Analysis")
    ws_cost_analysis.append(["Cost Analysis Summary"])
    ws_cost_analysis.append([])
    ws_cost_analysis.append(["Metric", "Value"])
    ws_cost_analysis.append(["Total Repairs", total_repairs])
    ws_cost_analysis.append(["Total Cost (KSh)", round(total_cost, 2)])
    ws_cost_analysis.append(["Labor Cost (KSh)", round(total_labor, 2)])
    ws_cost_analysis.append(["Parts Cost (KSh)", round(total_parts, 2)])
    ws_cost_analysis.append(["Additional Costs (KSh)", round(total_additional, 2)])
    ws_cost_analysis.append([])
    ws_cost_analysis.append(["Percentages"])
    if total_cost > 0:
        ws_cost_analysis.append(["Labor %", f"{(total_labor / total_cost * 100):.1f}%"])
        ws_cost_analysis.append(["Parts %", f"{(total_parts / total_cost * 100):.1f}%"])
        ws_cost_analysis.append(["Additional %", f"{(total_additional / total_cost * 100):.1f}%"])
    ws_cost_analysis.append([])
    ws_cost_analysis.append(["Averages"])
    if total_repairs > 0:
        ws_cost_analysis.append(["Avg Cost per Repair (KSh)", round(total_cost / total_repairs, 2)])
        ws_cost_analysis.append(["Avg Labor per Repair (KSh)", round(total_labor / total_repairs, 2)])
        ws_cost_analysis.append(["Avg Parts per Repair (KSh)", round(total_parts / total_repairs, 2)])
        ws_cost_analysis.append(["Avg Downtime per Repair (hrs)", round(total_downtime / total_repairs, 2)])

    # Calibration History Sheet (unchanged)
    ws_calibrations = wb.create_sheet("Calibration History")
    ws_calibrations.append([
        "Date", "Certificate Number", "Procedure", "Performed By",
        "Result", "Next Due Date", "Temperature", "Humidity", "Pressure"
    ])

    calibrations = CalibrationSession.objects.filter(
        device_serial=equipment.serial_number
    ).select_related('procedure', 'performed_by')

    for cal in calibrations:
        ws_calibrations.append([
            fmt_eat(cal.timestamp),
            cal.certificate_number,
            cal.procedure.name,
            cal.performed_by.get_full_name(),
            "PASS" if cal.overall_pass else "FAIL",
            cal.next_calibration_due.strftime('%Y-%m-%d') if cal.next_calibration_due else "N/A",
            f"{cal.actual_temperature}°C" if cal.actual_temperature else "N/A",
            f"{cal.actual_humidity}%RH" if cal.actual_humidity else "N/A",
            f"{cal.actual_pressure}kPa" if cal.actual_pressure else "N/A"
        ])

    # Checklist sheets: same data as the Machine Reports history modal
    from .repairs import checklist_history
    work_history, task_last_done = checklist_history(equipment)

    ws_tasks = wb.create_sheet("Checklist Tasks Last Done")
    ws_tasks.append(["Task", "Checklist", "Last Done", "Action", "Result", "Reading", "Note", "Work Order"])
    for t in task_last_done:
        ws_tasks.append([t['task'], t['checklist'], t['date'], t['action'], t['result'],
                         t['value'], t['note'], t['work_order']])

    ws_checklist = wb.create_sheet("Checklist History")
    ws_checklist.append(["Date", "Work Order", "Action", "Performed By", "Checklist", "Task",
                         "Expected", "Result", "Reading", "Note"])
    cards = jobcard.objects.filter(equipment=equipment, status='Approved', active_status=True) \
        .select_related('performed_by').prefetch_related('checklist_entries').order_by('-date_issued')
    for card in cards:
        for e in card.checklist_entries.all():
            ws_checklist.append([
                card.date_issued.strftime('%Y-%m-%d'), str(card.id)[:8].upper(), card.action_taken,
                card.performed_by.get_full_name() if card.performed_by else "N/A",
                e.template_title, e.task, e.expected_result, e.get_result_display(), e.value, e.note,
            ])

    # Prepare response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=equipment_history_{equipment.serial_number}_with_costs.xlsx'

    # Save the workbook to the response
    wb.save(response)

    return response


@login_required
def export_equipment_category_detailed_pdf(request):
    """
    Export equipment report to PDF with professional formatting

    Features:
    - Summary statistics with status breakdown
    - Category-based grouping
    - Workshop and category filtering
    - Search functionality
    - Professional header/footer with watermark
    """
    # Get user context
    profile, workshops, selected_workshop, is_hod = get_user_workshop_context(request)

    if profile is None:
        return HttpResponse("Access denied - No user profile assigned", status=403)

    # ==========================================
    # BUILD BASE EQUIPMENT QUERYSET
    # ==========================================
    equipment_qs = Equipment.objects.select_related(
        "workshop", "category", "description", "manufacturer"
    ).filter(active_status=True)

    # ==========================================
    # APPLY WORKSHOP FILTER
    # ==========================================
    if is_hod:
        # HOD can filter by any workshop
        workshop_id = request.GET.get("workshop")
        if workshop_id:
            try:
                equipment_qs = equipment_qs.filter(workshop_id=workshop_id)
                selected_workshop = Workshop.objects.get(id=workshop_id)
            except Workshop.DoesNotExist:
                pass
    else:
        # Regular users see only their workshop
        if selected_workshop:
            equipment_qs = equipment_qs.filter(workshop=selected_workshop)

    # ==========================================
    # APPLY CATEGORY FILTER
    # ==========================================
    category_id = request.GET.get("category")
    selected_category = None
    download_all_categories = request.GET.get("download_all_categories", "false") == "true"

    if category_id and not download_all_categories:
        try:
            selected_category = EquipmentCategory.objects.get(id=category_id)
            equipment_qs = equipment_qs.filter(category=selected_category)
        except EquipmentCategory.DoesNotExist:
            pass

    # ==========================================
    # APPLY SEARCH FILTER
    # ==========================================
    search_query = request.GET.get("search", "")
    if search_query:
        equipment_qs = equipment_qs.filter(
            Q(description__name__icontains=search_query) |
            Q(serial_number__icontains=search_query) |
            Q(model__icontains=search_query) |
            Q(manufacturer__name__icontains=search_query)
        )

    # ==========================================
    # DETERMINE REPORT TYPE
    # ==========================================
    # 'category' = show all categories grouped
    # 'detailed' = show single category or filtered results
    if download_all_categories or not selected_category:
        report_type = 'category'
    else:
        report_type = 'detailed'

    # ==========================================
    # GENERATE AND RETURN PDF
    # ==========================================
    return create_equipment_pdf_response(
        equipment_queryset=equipment_qs,
        selected_workshop=selected_workshop,
        selected_category=selected_category,
        search_query=search_query if search_query else None,
        report_type=report_type,
        generated_by=request.user
    )


class ManufacturerPerformanceDocTemplate(BaseDocTemplate):
    """Custom document template with headers and footers"""

    def __init__(self, filename, **kwargs):
        BaseDocTemplate.__init__(self, filename, **kwargs)

        # Define page templates
        frame = Frame(
            x1=2*cm, y1=2*cm,
            width=A4[0]-4*cm, height=A4[1]-5*cm,
            leftPadding=0, bottomPadding=0,
            rightPadding=0, topPadding=0
        )

        template = PageTemplate(
            id='normal',
            frames=[frame],
            onPage=self.add_header_footer
        )

        self.addPageTemplates([template])

    def add_header_footer(self, canvas, doc):
        """Add header and footer to each page"""
        canvas.saveState()

        # Header
        self.draw_header(canvas, doc)

        # Footer
        self.draw_footer(canvas, doc)

        canvas.restoreState()

    def draw_header(self, canvas, doc):
        """Draw header with logo and title"""
        # Header background (light gray bar)
        canvas.setFillColor(colors.lightgrey)
        canvas.rect(0, A4[1]-3*cm, A4[0], 3*cm, fill=1, stroke=0)

        # Logo (if exists)
        logo_path = os.path.join(settings.STATIC_ROOT or settings.STATICFILES_DIRS[0], 'images', 'logo.png')
        if os.path.exists(logo_path):
            try:
                canvas.drawImage(logo_path, 1*cm, A4[1]-2.5*cm, width=2*cm, height=1.5*cm, mask='auto')
            except Exception:
                # If logo fails to load, draw a placeholder
                canvas.setFillColor(colors.white)
                canvas.rect(1*cm, A4[1]-2.5*cm, 2*cm, 1.5*cm, fill=1, stroke=1)
                canvas.setFillColor(colors.black)
                canvas.setFont("Helvetica", 8)
                canvas.drawCentredString(2*cm, A4[1]-1.8*cm, "LOGO")
        else:
            # Draw placeholder logo box
            canvas.setFillColor(colors.white)
            canvas.rect(1*cm, A4[1]-2.5*cm, 2*cm, 1.5*cm, fill=1, stroke=1)
            canvas.setFillColor(colors.black)
            canvas.setFont("Helvetica", 8)
            canvas.drawCentredString(2*cm, A4[1]-1.8*cm, "LOGO")

        # Header title
        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawCentredString(A4[0]/2, A4[1]-1.2*cm, "MANUFACTURER PERFORMANCE REPORT")

        # Department info
        canvas.setFont("Helvetica", 10)
        canvas.drawCentredString(A4[0]/2, A4[1]-1.6*cm, "Biomedical Engineering Department")

        # Contact info
        canvas.setFont("Helvetica", 7)
        contact_text = contact_line("ISO/IEC 17025:2017")
        canvas.drawCentredString(A4[0]/2, A4[1]-2.3*cm, contact_text)

    def draw_footer(self, canvas, doc):
        """Draw footer with page number and additional info"""
        # Footer line
        canvas.setStrokeColor(colors.lightgrey)
        canvas.line(1*cm, 1.5*cm, A4[0]-1*cm, 1.5*cm)

        # Page number (lighter and italic)
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica-Oblique", 8)
        page_text = f"Page {doc.page}"
        canvas.drawRightString(A4[0]-1*cm, 1*cm, page_text)

        # Footer left text (lighter and italic)
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawString(1*cm, 1*cm, "Manufacturer Performance Analysis Report")

        # Footer center text (lighter and italic - watermark effect)
        canvas.setFillColor(colors.lightgrey)
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawCentredString(A4[0]/2, 1*cm, "Confidential - Internal Use Only")


@login_required
def export_manufacturer_performance_pdf(request):
    """
    Export manufacturer performance report to PDF

    Features:
    - Executive summary with key metrics
    - Performance overview table sorted by uptime
    - Detailed analysis for top manufacturers
    - Performance ratings (Excellent/Good/Fair/Poor)
    - Color-coded metrics
    - Recommendations based on data
    - Professional header/footer with watermark
    """
    # Get user context
    profile, workshops, selected_workshop, is_hod = get_user_workshop_context(request)

    if profile is None:
        return HttpResponse("Access denied - No user profile assigned", status=403)

    # ==========================================
    # BUILD BASE EQUIPMENT QUERYSET
    # ==========================================
    equipment_qs = Equipment.objects.select_related(
        "workshop", "category", "description", "manufacturer"
    ).filter(active_status=True)

    # ==========================================
    # APPLY WORKSHOP FILTER
    # ==========================================
    if is_hod:
        # HOD can filter by any workshop
        workshop_id = request.GET.get("workshop")
        if workshop_id:
            try:
                equipment_qs = equipment_qs.filter(workshop_id=workshop_id)
                selected_workshop = Workshop.objects.get(id=workshop_id)
            except Workshop.DoesNotExist:
                pass
    else:
        # Regular users see only their workshop
        if selected_workshop:
            equipment_qs = equipment_qs.filter(workshop=selected_workshop)

    # ==========================================
    # CALCULATE MANUFACTURER PERFORMANCE
    # ==========================================
    manufacturer_performance = calculate_manufacturer_performance(equipment_qs)

    # ==========================================
    # GENERATE AND RETURN PDF
    # ==========================================
    return create_manufacturer_pdf_response(
        manufacturer_performance_data=manufacturer_performance,
        selected_workshop=selected_workshop,
        generated_by=request.user
    )


class EquipmentListDocTemplate(BaseDocTemplate):
    """Custom document template for equipment list with headers and footers"""

    def __init__(self, filename, **kwargs):
        BaseDocTemplate.__init__(self, filename, **kwargs)

        # Define page templates - using landscape for better table display
        frame = Frame(
            x1=2*cm, y1=2*cm,
            width=A4[1]-4*cm, height=A4[0]-5*cm,
            leftPadding=0, bottomPadding=0,
            rightPadding=0, topPadding=0
        )

        template = PageTemplate(
            id='normal',
            frames=[frame],
            onPage=self.add_header_footer
        )

        self.addPageTemplates([template])

    def add_header_footer(self, canvas, doc):
        """Add header and footer to each page"""
        canvas.saveState()

        # Header
        self.draw_header(canvas, doc)

        # Footer
        self.draw_footer(canvas, doc)

        canvas.restoreState()

    def draw_header(self, canvas, doc):
        """Draw header with logo and title"""
        # Adjust for landscape orientation
        width, height = A4[1], A4[0]

        # Header background
        canvas.setFillColor(colors.lightgrey)
        canvas.rect(0, height-3*cm, width, 3*cm, fill=1, stroke=0)

        # Logo placeholder
        canvas.setFillColor(colors.white)
        canvas.rect(1*cm, height-2.5*cm, 2*cm, 1.5*cm, fill=1, stroke=1)
        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(2*cm, height-1.8*cm, "LOGO")

        # Header title
        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawCentredString(width/2, height-1.2*cm, "EQUIPMENT REPORT")

        # Department info
        canvas.setFont("Helvetica", 10)
        canvas.drawCentredString(width/2, height-1.6*cm, "Biomedical Engineering Department")

        # Contact info
        canvas.setFont("Helvetica", 7)
        contact_text = contact_line("ISO/IEC 17025:2017")
        canvas.drawCentredString(width/2, height-2.3*cm, contact_text)

    def draw_footer(self, canvas, doc):
        """Draw footer with page number and additional info"""
        width, height = A4[1], A4[0]

        # Footer line
        canvas.setStrokeColor(colors.lightgrey)
        canvas.line(1*cm, 1.5*cm, width-1*cm, 1.5*cm)

        # Page number
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica-Oblique", 8)
        page_text = f"Page {doc.page}"
        canvas.drawRightString(width-1*cm, 1*cm, page_text)

        # Footer left text
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawString(1*cm, 1*cm, "Equipment Report")

        # Footer center text
        canvas.setFillColor(colors.lightgrey)
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawCentredString(width/2, 1*cm, "Confidential - Internal Use Only")
